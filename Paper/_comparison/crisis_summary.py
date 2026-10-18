'''
Compact crisis summary from the two big crisis workbooks per frequency.

Inputs (written by crisis_analysis_defined.py and crisis_analysis_variable.py):
  _output/_comparison/crisis_analysis_all_<frequency>.xlsx           fix
  _output/_comparison/crisis_analysis_variable_all_<frequency>.xlsx  flex
for monthly, quarterly and yearly; fix and flex are joined on run (both must
hold the same runs).  Nothing is recomputed: every value is read from them.

Output: _output/_comparison/crisis_summary.xlsx
  <freq>_crises          one row per run, raw values per crisis C1 ... C6,
                         grouped by metric (full fix, full flex, ptt, ttp)
  overall                full sample 1998-2025, all frequencies stacked
  <freq>_crises_vs_mcap  the same against mcap of the same frequency and crisis
  overall_vs_mcap        the same for overall
  notes                  definitions
Rows: mcap, equal_weight, then per model (mvo, hrp, xgb, rf, lstm) the
baseline, its control (always min_var) and its regime runs (changepoint,
hmm, wasserstein; cash before the alternative target).  Columns of each
metric: C1 ... C6.  Days where the run is not recovered (flex) are lower
bounds: numeric, filled light orange.

vs_mcap: returns relative, (1 + r) / (1 + r_mcap) - 1 (checked against the
*_excess_return columns); drawdowns, Ulcer, ret_to_even and ratios as
differences; days as differences, empty if either side is censored
(C{i}_cens_vs_mcap = True); booleans and dates unchanged; mcap row = 0.

Usage: run the two crisis scripts first, then this file.
'''

import os
from pathlib import Path

import numpy as np
import pandas as pd

import compare_all as ca
import crisis_analysis_defined as cad
from crises import overall, main_crises
from metrics import annualized_return, sharpe_ratio


# ----
# Variables
# ----

FREQS     = ['monthly', 'quarterly', 'yearly']
CRISES    = [f'C{i}' for i in range(1, 7)]
SAVE_PATH = cad.SAVE_DIR / 'crisis_summary.xlsx'

ID_COLS     = ['model', 'run', 'frequency', 'role', 'detector', 'series', 'crisis_target',
               'base_run', 'fallback_run']
MODEL_ORDER = ['market_cap', 'equal_weight', 'mvo', 'hrp', 'xgb', 'rf', 'lstm']
ROLE_ORDER  = {'benchmark': 0, 'baseline': 0, 'control': 1, 'regime': 2}
DET_ORDER   = ['changepoint', 'hmm', 'wasserstein']
MCAP        = 'market_cap'

# crisis metrics: (name, source file, source column, group, kind)
#   kind: ret | dd | days | ulcer | diff | bool
CRISIS_METRICS = [
    ('ret_full_fix',         'fix',  'full_cum_return',    'full fix',  'ret'),
    ('dd_full_fix',          'fix',  'full_max_dd',        'full fix',  'dd'),
    ('rec_full_fix',         'fix',  'full_recovered',     'full fix',  'bool'),
    ('dd_full_flex',         'flex', 'full_anchored_dd',   'full flex', 'dd'),
    ('days_full_flex',       'flex', 'full_days',          'full flex', 'days'),
    ('rec_full_flex',        'flex', 'full_not_recovered', 'full flex', 'bool_not'),
    ('ulcer_flex',           'flex', 'full_ulcer_common',  'full flex', 'ulcer'),
    ('nodd_flex',            'flex', 'full_no_drawdown',   'full flex', 'bool'),
    ('ret_ptt_fix',          'fix',  'ptt_cum_return',     'ptt',       'ret'),
    ('dd_ptt_fix',           'fix',  'ptt_max_dd',         'ptt',       'dd'),
    ('days_ptt_flex',        'flex', 'ptt_days',           'ptt',       'days'),
    ('ret_ttp_fix',          'fix',  'ttp_cum_return',     'ttp',       'ret'),
    ('dd_ttp_fix',           'fix',  'ttp_max_dd',         'ttp',       'dd'),
    ('days_ttp_flex',        'flex', 'ttp_days',           'ttp',       'days'),
    ('ret_to_even_ttp_flex', 'flex', 'ttp_return_to_even', 'ttp',       'diff'),
]
CENSORED_DAYS = ('days_full_flex', 'days_ttp_flex')         # lower bounds when not recovered
EXCESS = {'ret_full_fix': 'full_excess_return', 'ret_ptt_fix': 'ptt_excess_return',
          'ret_ttp_fix': 'ttp_excess_return'}

# overall: (column, source file, source column, group, kind)
OVERALL_METRICS = (
    [(c, 'fix', c, 'full fix', k) for c, k in [
        ('full_cum_return', 'ret'), ('full_cagr', 'diff_pct'), ('full_ann_vol', 'diff_pct'),
        ('full_sharpe', 'diff'), ('full_sortino', 'diff'), ('full_calmar', 'diff'),
        ('full_max_dd', 'dd'), ('full_ulcer', 'ulcer'), ('full_turnover_pa', 'diff_pct')]]
    # computed here from the runs (not in the crisis workbooks), see calc_overall
    + [('ret_outside_crises_ann', 'calc', None, 'outside crises', 'diff_pct'),
       ('sharpe_outside_crises',  'calc', None, 'outside crises', 'diff'),
       ('sharpe_inside_crises',   'calc', None, 'outside crises', 'diff'),
       ('time_defensive_share',   'calc', None, 'outside crises', 'keep_pct')]
    + [(f'worst_{c}', 'flex', c, 'worst episode (flex)', k) for c, k in [
        ('peak_date', 'keep'), ('trough_date', 'keep'), ('recovery_date', 'keep'),
        ('full_anchored_dd', 'dd'), ('full_days', 'days'), ('full_not_recovered', 'bool')]]
    + [(f'{src}_hit_count_{p}', src, f'hit_count_{p}', 'hit counts vs mcap', 'keep')
       for src in ('fix', 'flex') for p in ('full', 'ptt', 'ttp')]
)

GROUP_FILL = {'full fix': 'DDEBF7', 'full flex': 'E2EFDA', 'ptt': 'E4DFEC', 'ttp': 'EDEDED',
              'outside crises': 'DDEBF7', 'worst episode (flex)': 'E2EFDA',
              'hit counts vs mcap': 'FFF2CC'}
CENSOR_FILL = 'F8CBAD'
PCT_KINDS = {'ret', 'dd', 'ulcer', 'diff', 'diff_pct', 'keep_pct'}


# ----
# Reading
# ----

def load(freq: str) -> dict:
    '''{'fix': {sheet: df}, 'flex': {sheet: df}} for one frequency, same runs asserted.'''
    out = {
        'fix':  pd.read_excel(cad.SAVE_DIR / f'crisis_analysis_all_{freq}.xlsx', sheet_name=None),
        'flex': pd.read_excel(cad.SAVE_DIR / f'crisis_analysis_variable_all_{freq}.xlsx', sheet_name=None),
    }
    for sheet in ['overall'] + CRISES:
        a, b = set(out['fix'][sheet]['run']), set(out['flex'][sheet]['run'])
        assert a == b, f'{freq} {sheet}: fix and flex differ in runs: {sorted(a ^ b)}'
    return out


def _order(row) -> tuple:
    m = row['model']
    target = str(row['crisis_target'])
    det = row['detector']
    return (MODEL_ORDER.index(m) if m in MODEL_ORDER else len(MODEL_ORDER), m,
            ROLE_ORDER.get(row['role'], 3),
            DET_ORDER.index(det) if det in DET_ORDER else -1,
            0 if target == 'cash' else 1, row['run'])


def ids(df: pd.DataFrame) -> pd.DataFrame:
    '''Identifier columns, rows in the summary order.'''
    out = df[ID_COLS].copy()
    out['_key'] = out.apply(_order, axis=1)
    return out.sort_values('_key').drop(columns='_key').reset_index(drop=True)


def _bool(s: pd.Series, negate: bool = False) -> pd.Series:
    '''True / False (negated if asked), None for empty cells.'''
    def conv(v):
        if pd.isna(v):
            return None
        return not bool(v) if negate else bool(v)
    return s.map(conv).astype(object)


def _is_false(v) -> bool:
    return v is not None and not pd.isna(v) and not bool(v)


# ----
# Crisis sheets
# ----

def crisis_raw(src: dict) -> tuple[pd.DataFrame, pd.DataFrame]:
    '''(raw sheet, censored mask with the same columns: True = lower bound).'''
    out = ids(src['fix']['C1'])
    cens = pd.DataFrame(False, index=out.index, columns=[])
    cols = {}
    for name, f, col, _, kind in CRISIS_METRICS:
        for c in CRISES:
            s = src[f][c].set_index('run')[col].reindex(out['run']).reset_index(drop=True)
            if kind in ('bool', 'bool_not'):
                s = _bool(s, negate=kind == 'bool_not')
            cols[f'{c}_{name}'] = s
    raw = pd.concat([out, pd.DataFrame(cols)], axis=1)
    for c in CRISES:
        not_rec = raw[f'{c}_rec_full_flex'].eq(False)
        for name in CENSORED_DAYS:
            cens[f'{c}_{name}'] = not_rec
    return raw, cens


def crisis_vs_mcap(raw: pd.DataFrame, src: dict) -> pd.DataFrame:
    '''raw against the mcap row; returns asserted against the *_excess_return columns.'''
    m = raw[raw['model'] == MCAP].iloc[0]
    out = raw[ID_COLS].copy()
    cols = {}
    for name, f, col, _, kind in CRISIS_METRICS:
        for c in CRISES:
            k, x = f'{c}_{name}', raw[f'{c}_{name}']
            if kind == 'ret':
                cols[k] = (1 + x) / (1 + m[k]) - 1
                exc = src['fix'][c].set_index('run')[EXCESS[name]].reindex(raw['run']).to_numpy()
                assert np.allclose(cols[k].to_numpy(dtype=float), exc, equal_nan=True, atol=1e-12), \
                    f'{k}: relative excess differs from {EXCESS[name]}'
            elif kind in ('bool', 'bool_not'):
                cols[k] = x
            elif name in CENSORED_DAYS:
                cens = raw[f'{c}_rec_full_flex'].eq(False) | _is_false(m[f'{c}_rec_full_flex'])
                cols[k] = (x - m[k]).where(~cens)
                if name == 'days_full_flex':
                    cols[f'{c}_cens_vs_mcap'] = cens
            else:
                cols[k] = x - m[k]
    vs = pd.concat([out, pd.DataFrame(cols)], axis=1)
    # column order: as raw, the censoring flags after the days_full_flex block
    order = list(ID_COLS)
    for name, *_ in CRISIS_METRICS:
        order += [f'{c}_{name}' for c in CRISES]
        if name == 'days_full_flex':
            order += [f'{c}_cens_vs_mcap' for c in CRISES]
    return vs[order]


# ----
# Overall sheets
# ----

def calc_overall(out: pd.DataFrame, runs: dict) -> dict:
    '''
    Columns computed from the runs, over the overall window (1998 .. 2025):
      ret_outside_crises_ann  annualised return (metrics.annualized_return) on
                              the trading days outside every main crisis
                              window (S&P peak, break-even]
      sharpe_outside_crises   annualised Sharpe ratio (metrics.sharpe_ratio,
      sharpe_inside_crises    same risk-free rate) of the pooled daily returns
                              outside / inside the main crisis windows
      time_defensive_share    regime runs: share of trading days in the
                              crisis portfolio (crisis_analysis_defined.defensive:
                              decided on rebalancing dates, effective the next
                              day); empty for runs without a detector
    '''
    w = overall[0]
    ret_out, sr_out, sr_in, defn = [], [], [], []
    for run, freq, role in zip(out['run'], out['frequency'], out['role']):
        r = runs[(run, freq)]
        ret = ca._slice(r['returns'], w['peak'], w['even'])
        inside = np.zeros(len(ret), dtype=bool)
        for c in main_crises:
            inside |= (ret.index > pd.Timestamp(c['peak'])) & (ret.index <= pd.Timestamp(c['even']))
        ret_out.append(annualized_return(ret[~inside]))
        sr_out.append(sharpe_ratio(ret[~inside]))
        sr_in.append(sharpe_ratio(ret[inside]))
        d = cad.defensive(r) if role == 'regime' else None
        defn.append(float(d.reindex(ret.index, fill_value=False).mean()) if d is not None else np.nan)
    return {'ret_outside_crises_ann': ret_out, 'sharpe_outside_crises': sr_out,
            'sharpe_inside_crises': sr_in, 'time_defensive_share': defn}


def overall_raw(srcs: dict, runs: dict) -> tuple[pd.DataFrame, list[str]]:
    '''(overall sheet with all frequencies, skipped columns).'''
    parts, skipped = [], set()
    for freq in FREQS:
        src = srcs[freq]
        out = ids(src['fix']['overall'])
        calc = calc_overall(out, runs)
        for name, f, col, _, kind in OVERALL_METRICS:
            if f == 'calc':
                out[name] = calc[name]
                continue
            tab = src[f]['overall'].set_index('run')
            if col not in tab:
                skipped.add(f'{f}:{col}')
                continue
            s = tab[col].reindex(out['run']).reset_index(drop=True)
            out[name] = _bool(s) if kind == 'bool' else s
        parts.append(out)
    return pd.concat(parts, ignore_index=True), sorted(skipped)


def overall_vs_mcap(raw: pd.DataFrame) -> pd.DataFrame:
    out = raw.copy()
    for freq in FREQS:
        rows = raw['frequency'] == freq
        m = raw[rows & (raw['model'] == MCAP)].iloc[0]
        for name, f, col, _, kind in OVERALL_METRICS:
            if name not in raw:
                continue
            x = raw.loc[rows, name]
            if kind == 'ret':
                out.loc[rows, name] = (1 + x) / (1 + m[name]) - 1
            elif kind == 'days':
                cens = raw.loc[rows, 'worst_full_not_recovered'].eq(True) | bool(m['worst_full_not_recovered'])
                out.loc[rows, name] = (x - m[name]).where(~cens)
            elif kind in ('dd', 'ulcer', 'diff', 'diff_pct'):
                out.loc[rows, name] = x - m[name]
    return out


# ----
# Writing
# ----

def _kind_of(col: str) -> tuple[str | None, str | None]:
    '''(group, kind) of an output column.'''
    for name, _, _, group, kind in CRISIS_METRICS:
        if col.endswith('_' + name) and col[:2] in {c[:2] for c in CRISES}:
            return group, kind
    if col.endswith('_cens_vs_mcap'):
        return 'full flex', 'bool'
    for name, _, _, group, kind in OVERALL_METRICS:
        if col == name:
            return group, kind
    return None, None


def _format(ws, df: pd.DataFrame, cens: pd.DataFrame | None) -> None:
    from openpyxl.styles import Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    n_ids = len(ID_COLS)
    ws.freeze_panes = get_column_letter(n_ids + 1) + '2'
    ws.auto_filter.ref = ws.dimensions
    top = Border(top=Side(style='thin', color='808080'))
    orange = PatternFill('solid', fgColor=CENSOR_FILL)
    for j, col in enumerate(df.columns, 1):
        letter = get_column_letter(j)
        group, kind = _kind_of(col)
        head = ws.cell(row=1, column=j)
        head.font = Font(bold=True)
        width = len(str(col)) + 2
        ws.column_dimensions[letter].width = min(max(9, width), 30 if j <= n_ids else 22)
        fill = PatternFill('solid', fgColor=GROUP_FILL[group]) if group in GROUP_FILL else None
        if fill is not None:
            head.fill = fill
        is_ratio = any(x in col for x in ('sharpe', 'sortino', 'calmar'))
        fmt = ('0.00%' if kind in PCT_KINDS and not is_ratio
               else '0' if kind == 'days'
               else 'yyyy-mm-dd' if col.endswith('_date')
               else '0.000' if kind == 'diff' else None)
        for i in range(len(df)):
            cell = ws.cell(row=i + 2, column=j)
            if fmt:
                cell.number_format = fmt
            if fill is not None:
                cell.fill = fill
            if cens is not None and col in cens and bool(cens[col].iloc[i]):
                cell.fill = orange
    # thin line where the model (or, in overall, the frequency) changes
    key = df['frequency'].astype(str) + '|' + df['model'].astype(str)
    for i in range(1, len(df)):
        if key.iloc[i] != key.iloc[i - 1]:
            for j in range(1, len(df.columns) + 1):
                ws.cell(row=i + 2, column=j).border = top


NOTES = [
    'fix = pre-defined S&P 500 windows (crises.py): every window starts the day after the S&P peak with '
    'the portfolio value set to 1; full = peak -> break-even, ptt = peak -> trough, ttp = trough -> '
    'break-even (crisis_analysis_defined.py).',
    'flex = common start at the S&P peak, then each portfolio\'s own trough and recovery (back at its start '
    'value), recovery searched up to the next crisis peak (crisis_analysis_variable.py).',
    'ulcer_flex = Ulcer index over a common horizon per crisis, from the crisis peak to the next crisis '
    'peak (C6: end of the data, 2025-12-31), drawdown from the running high until the own recovery, 0 after; '
    'the same number of days for every portfolio, so a longer recovery raises it.',
    'Censoring: where rec_full_flex is False the portfolio had not recovered by the next crisis peak (C6: by '
    'the end of the data); days_full_flex and days_ttp_flex then count up to that cut-off and are lower '
    'bounds ("at least"). These cells are numeric and filled light orange.',
    'dd_full_flex = anchored drawdown (lowest value / start value - 1, capped at 0); dd_*_fix = maximum '
    'drawdown from the running high within the window.',
    'Crisis windows are dated on the S&P 500 price index; all portfolios and mcap are total-return series.',
    'vs_mcap sheets, against mcap at the same frequency and crisis: returns relative, (1 + r) / (1 + r_mcap) '
    '- 1 (equal to the *_excess_return columns of the source); drawdowns dd - dd_mcap (positive = shallower); '
    'days days - days_mcap (negative = faster), empty if either side is censored (days_full_flex and '
    'days_ttp_flex; C{i}_cens_vs_mcap = True); Ulcer ulcer - ulcer_mcap (negative = less pain); '
    'ret_to_even as a difference; booleans unchanged; the mcap row is 0.',
    'ret_outside_crises_ann (overall): annualised return on the trading days 1998-2025 outside all six '
    'crisis windows (S&P peak -> break-even); computed from the daily returns of each run.',
    'sharpe_outside_crises / sharpe_inside_crises (overall): annualised Sharpe ratio (same risk-free rate as '
    'full_sharpe) of the pooled daily returns outside / inside the six crisis windows (inside: all six '
    'crises together, S&P peak -> break-even).',
    'time_defensive_share (overall): share of trading days 1998-2025 in the crisis portfolio, decided on '
    'the rebalancing dates and effective from the next day; regime runs only.',
    'overall_vs_mcap: full_cum_return relative; CAGR, vol, turnover, ratios, drawdown, Ulcer and '
    'ret_outside_crises_ann as differences; time_defensive_share unchanged; '
    'worst-episode days as a difference (empty if either is not recovered); dates and hit counts unchanged.',
    'Hit counts (overall): crises won against mcap, "wins/comparable crises" (fix: higher return in the '
    'phase; flex: full = lower common Ulcer, ptt = shallower trough, ttp = faster recovery).',
    'Detectors use different input series: changepoint = VIX, HMM = S&P 500 (GSPC), Wasserstein = S&P 500 '
    '+ VIX; differences between detectors reflect method plus data.',
    'A detector can switch only on the rebalancing dates of its frequency (yearly: at most once a year).',
]


def write(sheets: dict, cens: dict, path: Path) -> bool:
    '''Write via a temp file and rename; False (with a message) if the target is locked.'''
    tmp = path.with_name(path.stem + '.tmp.xlsx')
    with pd.ExcelWriter(tmp, engine='openpyxl') as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
            if name == 'notes':
                xw.sheets[name].column_dimensions['A'].width = 140
            else:
                _format(xw.sheets[name], df, cens.get(name))
    try:
        os.replace(tmp, path)
    except PermissionError:
        print(f'[crisis_summary] {path.name} is open in Excel (locked): close it and run again; '
              f'the new version is in {tmp.name}')
        return False
    print(f'[crisis_summary] saved to {path}')
    return True


def main() -> None:
    srcs = {f: load(f) for f in FREQS}
    sheets, cens = {}, {}
    vs = {}
    for f in FREQS:
        raw, mask = crisis_raw(srcs[f])
        sheets[f'{f}_crises'], cens[f'{f}_crises'] = raw, mask
        vs[f'{f}_crises_vs_mcap'] = crisis_vs_mcap(raw, srcs[f])
    runs = {(r['run'], r['frequency']): r for r in ca.find_runs()}
    ov, skipped = overall_raw(srcs, runs)
    sheets['overall'] = ov
    sheets.update(vs)
    sheets['overall_vs_mcap'] = overall_vs_mcap(ov)
    sheets['notes'] = pd.DataFrame({'note': NOTES})
    print('[crisis_summary] vs_mcap returns match the *_excess_return columns')
    if skipped:
        print(f'[crisis_summary] source columns not found, left out: {", ".join(skipped)}')
    for name, df in sheets.items():
        if name != 'notes':
            print(f'[crisis_summary] {name}: {len(df)} rows, {len(df.columns)} columns')
    write(sheets, cens, SAVE_PATH)


if __name__ == '__main__':
    main()
