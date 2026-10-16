'''
Is the ranking of the regime strategies stable across crises?  Friedman test
with Nemenyi post-hoc test and critical difference diagram (Demsar 2006).

Methods (the k columns): the baseline and every detector x crisis target,
e.g. 'base', 'changepoint cash', 'changepoint defensive', ... (defensive = hrp
for the ML models, min_var for MVO and HRP).
Blocks (the N rows): one per model, frequency and crisis.  For phase PTT the
metric over peak -> trough, for TTP over trough -> recovery, both for all
main crises (N = models x frequencies x crises); for Overall the full period
(N = models x frequencies).  Within each block the k methods share model and
frequency, so they are ranked against each other (1 = best, ties averaged).

Sheets of friedman.xlsx (in _output/_comparison/):
  summary    per phase and metric: N, k, Friedman chi2 and p, Iman-Davenport
             F and p, Kendall's W (0 = no agreement between blocks on the
             order, 1 = the same order in every block), Nemenyi critical
             difference CD at ALPHA
  avg_ranks  average rank of every method per phase and metric (lower is
             better)
  pairwise   every pair of methods: average rank difference, Nemenyi p,
             Wilcoxon signed-rank p with Holm correction (Benavoli et al.
             2016), significant = Nemenyi p < ALPHA
  by_crisis  average rank of every method within each crisis (blocks =
             models x frequencies), and per phase and metric Kendall's W of
             the crises: do the crises agree on the order of the methods?
CD diagrams: cd_<phase>_<metric>.png (needs matplotlib): methods on the
average-rank axis, methods joined by a bar are not significantly different.

Metric values are computed with compare_all.window_metrics, as in
compare_all.xlsx and the zscores workbooks.

Usage: set the variables below and run the file.
'''

import numpy as np
import pandas as pd
from scipy import stats

import compare_all as ca
import z_scores as zs
from compare_runs import CRISIS_NAMES
from crises import main_crises

# ----
# Variables
# ----

# metric -> better direction (+1 higher is better)
METRICS = {
    'cum_return': +1,        # return over the window: the loss avoided (PTT) / gain kept (TTP)
    'max_dd':     +1,
    'sharpe':     +1,
    'calmar':     +1,
}
PHASES       = ['PTT', 'TTP', 'Overall']
MODELS       = ['mvo', 'hrp', 'xgb', 'rf', 'lstm']
TARGET_GROUPS = {'cash': 'cash', 'hrp': 'defensive', 'min_var': 'defensive'}
ALPHA        = 0.05
CD_PLOTS     = True

SAVE_PATH = ca.SAVE_DIR / 'friedman.xlsx'


# ----
# Data
# ----

def method(run: dict) -> str | None:
    if run['detector'] == 'none':
        return 'base' if run['crisis_target'] == '' else None       # skip the always-min-var control
    t = TARGET_GROUPS.get(run['crisis_target'])
    return run['detector'] + ' ' + t if t else None


def blocks(values: dict, runs: list[dict], phase: str, metric: str) -> pd.DataFrame:
    '''block x method table of the metric (block = model | frequency | crisis).'''
    v = values[metric]
    if phase == 'Overall':
        rows = {'Overall': 'Overall'}
    else:
        suffix = ' peak -> trough' if phase == 'PTT' else ' trough -> recovery'
        rows = {l: l[:-len(suffix)] for l in v.index if l.endswith(suffix)}
    recs = []
    for r in runs:
        m = method(r)
        if m is None or r['model'] not in MODELS:
            continue
        for label, crisis in rows.items():
            recs.append({'model': r['model'], 'frequency': r['frequency'], 'crisis': crisis,
                         'method': m, 'value': v.loc[label, r['column']]})
    df = pd.DataFrame(recs)
    if df.empty:
        return df
    t = df.pivot_table(index=['model', 'frequency', 'crisis'], columns='method', values='value',
                       aggfunc='first')
    order = ['base'] + sorted(c for c in t.columns if c != 'base')
    return t[[c for c in order if c in t.columns]].dropna()


# ----
# Tests
# ----

def ranks(table: pd.DataFrame, direction: int) -> pd.DataFrame:
    '''Rank within every block, 1 = best, ties averaged (inf counts as a value).'''
    return (-direction * table).rank(axis=1, method='average')


def friedman(rk: pd.DataFrame) -> dict:
    N, k = rk.shape
    R = rk.mean(axis=0).to_numpy()
    chi2 = 12 * N / (k * (k + 1)) * (np.sum(R ** 2) - k * (k + 1) ** 2 / 4)
    # tie correction (Conover): divide by 1 - sum(t^3 - t) / (N k (k^2 - 1))
    ties = sum(((c ** 3 - c).sum() for c in (row.value_counts().to_numpy() for _, row in rk.iterrows())))
    corr = 1 - ties / (N * k * (k ** 2 - 1))
    chi2 = chi2 / corr if corr > 0 else np.nan
    ff = (N - 1) * chi2 / (N * (k - 1) - chi2) if N * (k - 1) > chi2 else np.inf
    q = stats.studentized_range.ppf(1 - ALPHA, k, 1e6) / np.sqrt(2)
    return {'N': N, 'k': k, 'chi2': chi2, 'p_friedman': stats.chi2.sf(chi2, k - 1),
            'F_iman_davenport': ff, 'p_iman_davenport': stats.f.sf(ff, k - 1, (k - 1) * (N - 1)),
            'kendall_W': chi2 / (N * (k - 1)), 'CD': q * np.sqrt(k * (k + 1) / (6 * N))}


def _holm(p: np.ndarray) -> np.ndarray:
    p = np.asarray(p, float)
    out = np.full(len(p), np.nan)
    ok = np.flatnonzero(~np.isnan(p))
    order = ok[np.argsort(p[ok])]
    m = len(order)
    adj = np.maximum.accumulate([(m - i) * p[j] for i, j in enumerate(order)])
    out[order] = np.minimum(adj, 1.0)
    return out


def pairwise(table: pd.DataFrame, rk: pd.DataFrame, direction: int) -> pd.DataFrame:
    N, k = rk.shape
    R = rk.mean(axis=0)
    se = np.sqrt(k * (k + 1) / (6 * N))
    rows = []
    for i, a in enumerate(rk.columns):
        for b in rk.columns[i + 1:]:
            q = abs(R[a] - R[b]) / se
            x = direction * (table[a] - table[b]).replace([np.inf, -np.inf], np.nan).dropna()
            try:
                pw = stats.wilcoxon(x).pvalue if (x != 0).any() else 1.0
            except ValueError:
                pw = np.nan
            rows.append({'method a': a, 'method b': b, 'rank a': R[a], 'rank b': R[b],
                         'rank_diff': R[a] - R[b],
                         'p_nemenyi': min(1.0, stats.studentized_range.sf(q * np.sqrt(2), k, 1e6)),
                         'p_wilcoxon': pw})
    df = pd.DataFrame(rows)
    df['p_wilcoxon_holm'] = _holm(df['p_wilcoxon'].to_numpy())
    df['significant'] = df['p_nemenyi'] < ALPHA
    return df


def by_crisis(rk: pd.DataFrame) -> tuple[pd.DataFrame, float]:
    '''Average rank per crisis and Kendall's W of the crises.'''
    avg = rk.groupby(level='crisis', sort=False).mean()
    if len(avg) < 2:
        return avg, np.nan
    return avg, friedman(avg.rank(axis=1, method='average'))['kendall_W']


# ----
# Critical difference diagram
# ----

def cd_plot(avg: pd.Series, cd: float, title: str, path) -> None:
    try:
        import matplotlib
        matplotlib.use('Agg')
        import matplotlib.pyplot as plt
    except ImportError:
        print('[friedman] matplotlib not installed: no CD diagram')
        return
    avg = avg.sort_values()
    k = len(avg)
    fig, ax = plt.subplots(figsize=(8, 1.0 + 0.25 * k))
    ax.set_xlim(k + 0.3, 0.7)                        # rank 1 on the right, as in Demsar (2006)
    ax.axis('off')
    y0 = 0.9
    ax.hlines(y0, 1, k, color='black', lw=1)
    for r in range(1, k + 1):
        ax.vlines(r, y0, y0 + 0.03, color='black', lw=1)
        ax.text(r, y0 + 0.05, str(r), ha='center', va='bottom', fontsize=8)
    ax.hlines(y0 + 0.16, 1, 1 + cd, color='black', lw=2)
    ax.text(1 + cd / 2, y0 + 0.18, f'CD = {cd:.2f}', ha='center', va='bottom', fontsize=8)
    step = 0.12
    lowest = y0
    for i, (name, r) in enumerate(avg.items()):
        left = i >= k / 2                            # better half labelled on the right
        y = y0 - step * (i + 1 if not left else k - i)
        x_text = 0.75 if not left else k + 0.25
        lowest = min(lowest, y)
        ax.vlines(r, y, y0, color='grey', lw=0.8)
        ax.hlines(y, r, x_text, color='grey', lw=0.8)
        ax.text(x_text, y, f'{name} ({r:.2f})', ha='left' if not left else 'right',
                va='center', fontsize=8)
    # cliques: maximal runs of methods within CD of each other
    vals = avg.to_numpy()
    cliques = []
    for i in range(k):
        j = i
        while j + 1 < k and vals[j + 1] - vals[i] <= cd:
            j += 1
        if j > i and not any(a <= i and j <= b for a, b in cliques):
            cliques.append((i, j))
    for n, (i, j) in enumerate(cliques):
        y = y0 - 0.04 - 0.035 * n
        ax.hlines(y, vals[i] - 0.03, vals[j] + 0.03, color='black', lw=3)
    ax.set_ylim(lowest - 0.08, 1.12)
    ax.set_title(title, fontsize=9)
    fig.tight_layout()
    fig.savefig(path, dpi=200)
    plt.close(fig)


# ----
# Main
# ----

def _write(sheets: dict) -> None:
    from openpyxl.utils import get_column_letter
    with pd.ExcelWriter(SAVE_PATH, engine='openpyxl') as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            ws.freeze_panes = 'C2'
            for j, col in enumerate(df.columns, 1):
                letter = get_column_letter(j)
                ws.column_dimensions[letter].width = min(max(10, len(str(col)) + 2), 40)
                if df[col].dtype.kind == 'f':
                    fmt = '0.0000' if str(col).startswith('p_') else '0.00'
                    for cell in ws[letter][1:]:
                        cell.number_format = fmt


def main() -> None:
    runs, benches = zs.select_runs()
    runs = [r for r in runs if r['model'] in MODELS and method(r) is not None]
    if not runs:
        print('[friedman] no runs found')
        return
    old = zs.METRICS
    zs.METRICS = {m: d for m, d in METRICS.items()}               # compute only what is needed
    print(f'[friedman] {len(runs)} runs; computing {len(METRICS)} metrics over {len(zs.periods())} periods')
    values = zs.value_tables(runs, benches)
    zs.METRICS = old

    summary, avg_rows, pair_rows, crisis_rows = [], [], [], []
    for phase in PHASES:
        for metric, direction in METRICS.items():
            table = blocks(values, runs, phase, metric)
            if table.empty or table.shape[1] < 2 or len(table) < 2:
                print(f'[friedman] {phase} {metric}: too few blocks or methods, skipped')
                continue
            rk = ranks(table, direction)
            res = friedman(rk)
            summary.append({'phase': phase, 'metric': metric, **res})
            avg = rk.mean(axis=0)
            avg_rows.append({'phase': phase, 'metric': metric, **avg.to_dict()})
            pw = pairwise(table, rk, direction)
            pw.insert(0, 'metric', metric)
            pw.insert(0, 'phase', phase)
            pair_rows.append(pw)
            if phase != 'Overall':
                per, w = by_crisis(rk)
                per = per.reset_index()
                per.insert(0, 'metric', metric)
                per.insert(0, 'phase', phase)
                per['kendall_W_crises'] = w
                crisis_rows.append(per)
            if CD_PLOTS:
                n, p = res['N'], res['p_friedman']
                cd_plot(avg, res['CD'], f'{phase} {metric}: N = {n}, Friedman p = {p:.3g}',
                        ca.SAVE_DIR / f'cd_{phase}_{metric}.png')
    sheets = {'summary': pd.DataFrame(summary), 'avg_ranks': pd.DataFrame(avg_rows)}
    if pair_rows:
        sheets['pairwise'] = pd.concat(pair_rows, ignore_index=True)
    if crisis_rows:
        sheets['by_crisis'] = pd.concat(crisis_rows, ignore_index=True)
    ca.SAVE_DIR.mkdir(parents=True, exist_ok=True)
    _write(sheets)
    cols = ['phase', 'metric', 'N', 'p_friedman', 'kendall_W', 'CD']
    for phase, metric, n, p, w, cd in sheets['summary'][cols].itertuples(index=False):
        print(f'[friedman] {phase:7s} {metric:10s} N={n:3d} p={p:.3g} W={w:.2f} CD={cd:.2f}')
    print(f'[friedman] saved to {SAVE_PATH}')


if __name__ == '__main__':
    main()
