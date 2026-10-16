'''
Which design choice drives the effect of the regime switch?

Every regime run is one cell of a full factorial design: model x rebalancing
frequency x detector x crisis target.  The response is the run's difference
to its own baseline (the d_<metric> columns of sheet vs_base in
compare_all.xlsx, so run compare_all.py first).  For each response a linear
model with the factors (and INTERACTIONS) as categorical regressors is fitted
and its variance is split over the factors (type II sums of squares).

Crisis target: 'cash' vs 'defensive' (hrp for the ML models, min_var for MVO
and HRP; TARGET_GROUPS), so that it is crossed with every model.

Sheets of factors.xlsx (in _output/_comparison/):
  anova    per response and term: df, sum of squares, eta2 = SS / SS_total
           (share of the variation across runs explained by the term),
           partial eta2 = SS / (SS + SS_residual), F and p.  The runs share
           baselines and market paths, so they are not independent draws:
           read eta2 as a descriptive split of the variation and the p-values
           as approximate.
  means    per response, factor and level: mean, median and standard
           deviation of the response, number of runs and share of runs that
           improve on their baseline (in the metric's better direction,
           BETTER below)
  cells    per response: mean response for every detector x target (rows)
           and frequency (columns)
  runs     the data used

Usage: set the variables below and run the file.
'''

import itertools

import numpy as np
import pandas as pd
from scipy import stats

import compare_all as ca

# ----
# Variables
# ----

SOURCE = ca.SAVE_DIR / 'compare_all.xlsx'

# response -> better direction of the difference (+1: a higher run - baseline is better)
BETTER = {
    'd_sharpe':     +1,
    'd_ann_return': +1,
    'd_ann_vol':    -1,
    'd_max_dd':     +1,      # max_dd is negative: a positive difference = smaller drawdown
    'd_calmar':     +1,
    'd_cvar_95':    +1,
    'd_ulcer':      -1,
}

FACTORS      = ['model', 'frequency', 'detector', 'target']
INTERACTIONS = [('detector', 'target'), ('detector', 'frequency'), ('model', 'target')]
TARGET_GROUPS = {'cash': 'cash', 'hrp': 'defensive', 'min_var': 'defensive'}
MODELS       = ['mvo', 'hrp', 'xgb', 'rf', 'lstm']

SAVE_PATH = ca.SAVE_DIR / 'factors.xlsx'


# ----
# Data
# ----

def load() -> pd.DataFrame:
    df = pd.read_excel(SOURCE, sheet_name='vs_base')
    df = df[df['detector'].isin(ca.DETECTORS) & df['model'].isin(MODELS)
            & df['crisis_target'].isin(TARGET_GROUPS)].copy()
    df['target'] = df['crisis_target'].map(TARGET_GROUPS)
    cells = df.groupby(FACTORS).size()
    full = np.prod([df[f].nunique() for f in FACTORS])
    if len(cells) < full or cells.max() > 1:
        print(f'[factors] note: {len(cells)} of {full} cells filled, up to {cells.max()} runs per '
              f'cell; the design is not balanced, type II sums of squares still apply')
    return df.reset_index(drop=True)


# ----
# Linear model and type II sums of squares
# ----

def _dummies(df: pd.DataFrame, factor: str) -> pd.DataFrame:
    '''Treatment coding: one column per level except the first.'''
    levels = sorted(df[factor].unique())
    return pd.DataFrame({f'{factor}={l}': (df[factor] == l).astype(float) for l in levels[1:]},
                        index=df.index)


def _columns(df: pd.DataFrame, term: tuple) -> pd.DataFrame:
    '''Design columns of a main effect (1 factor) or interaction (product of dummies).'''
    parts = [_dummies(df, f) for f in term]
    out = parts[0]
    for p in parts[1:]:
        out = pd.DataFrame({f'{a}:{b}': out[a] * p[b] for a in out for b in p}, index=df.index)
    return out


def _rss(y: np.ndarray, blocks: list[pd.DataFrame]) -> tuple[float, int]:
    X = np.column_stack([np.ones(len(y))] + [b.to_numpy() for b in blocks])
    b, *_ = np.linalg.lstsq(X, y, rcond=None)
    rank = np.linalg.matrix_rank(X)
    u = y - X @ b
    return float(u @ u), int(rank)


def anova(df: pd.DataFrame, response: str) -> pd.DataFrame:
    terms = [(f,) for f in FACTORS] + [tuple(t) for t in INTERACTIONS]
    d = df.dropna(subset=[response])
    y = d[response].to_numpy(float)
    cols = {t: _columns(d, t) for t in terms}
    rss_full, rank_full = _rss(y, list(cols.values()))
    df_res = len(y) - rank_full
    ss_tot = float(((y - y.mean()) ** 2).sum())
    rows = []
    for t in terms:
        higher = [h for h in terms if h != t and set(t) <= set(h)]     # terms containing t
        base = [cols[h] for h in terms if h != t and h not in higher]
        rss_without, rank_without = _rss(y, base)
        rss_with, rank_with = _rss(y, base + [cols[t]])
        ss, dof = rss_without - rss_with, rank_with - rank_without
        f = (ss / dof) / (rss_full / df_res) if dof and df_res and rss_full > 0 else np.nan
        rows.append({'response': response, 'term': ' x '.join(t), 'df': dof, 'SS': ss,
                     'eta2': ss / ss_tot if ss_tot else np.nan,
                     'partial_eta2': ss / (ss + rss_full) if ss + rss_full else np.nan,
                     'F': f, 'p': stats.f.sf(f, dof, df_res) if np.isfinite(f) else np.nan})
    rows.append({'response': response, 'term': 'Residual', 'df': df_res, 'SS': rss_full,
                 'eta2': rss_full / ss_tot if ss_tot else np.nan})
    rows.append({'response': response, 'term': 'Total', 'df': len(y) - 1, 'SS': ss_tot, 'eta2': 1.0})
    return pd.DataFrame(rows)


# ----
# Descriptive tables
# ----

def means(df: pd.DataFrame, response: str) -> pd.DataFrame:
    rows = []
    for f in FACTORS:
        for level, g in df.groupby(f):
            v = g[response].dropna()
            rows.append({'response': response, 'factor': f, 'level': level, 'runs': len(v),
                         'mean': v.mean(), 'median': v.median(), 'std': v.std(ddof=1),
                         'share_better': float((BETTER[response] * v > 0).mean()) if len(v) else np.nan})
    return pd.DataFrame(rows)


def cells(df: pd.DataFrame, response: str) -> pd.DataFrame:
    t = df.pivot_table(index=['detector', 'target'], columns='frequency', values=response, aggfunc='mean')
    t = t[[c for c in ca.FREQ_ORDER if c in t.columns]]
    t['all'] = df.groupby(['detector', 'target'])[response].mean()
    t = t.reset_index()
    t.insert(0, 'response', response)
    return t


# ----
# Main
# ----

def _write(sheets: dict) -> None:
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    with pd.ExcelWriter(SAVE_PATH, engine='openpyxl') as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            ws.freeze_panes = 'C2'
            for j, col in enumerate(df.columns, 1):
                letter = get_column_letter(j)
                ws.column_dimensions[letter].width = min(max(10, len(str(col)) + 2), 40)
                fmt = ('0.0%' if col in ('eta2', 'partial_eta2', 'share_better')
                       else '0.0000' if df[col].dtype.kind == 'f' else None)
                if fmt:
                    for cell in ws[letter][1:]:
                        cell.number_format = fmt
            if name == 'anova':                    # bold the largest eta2 of every response
                top = df[~df['term'].isin(['Residual', 'Total'])].groupby('response')['eta2'].idxmax()
                for i in top.dropna():
                    for c in ws[int(i) + 2]:
                        c.font = Font(bold=True)


def main() -> None:
    df = load()
    if df.empty:
        print(f'[factors] no regime runs in {SOURCE}')
        return
    responses = [r for r in BETTER if r in df.columns]
    names = ', '.join(responses)
    print(f'[factors] {len(df)} regime runs, responses: {names}')
    a = pd.concat([anova(df, r) for r in responses], ignore_index=True)
    sheets = {'anova': a,
              'means': pd.concat([means(df, r) for r in responses], ignore_index=True),
              'cells': pd.concat([cells(df, r) for r in responses], ignore_index=True),
              'runs': df[['model', 'frequency', 'detector', 'series', 'crisis_target', 'target', 'run']
                         + responses]}
    _write(sheets)
    top = a[~a['term'].isin(['Residual', 'Total'])].sort_values('eta2', ascending=False) \
        .groupby('response', sort=False).head(1)
    for resp, term, eta2 in top[['response', 'term', 'eta2']].itertuples(index=False):
        print(f'[factors] {resp}: largest share {term} ({eta2:.0%})')
    print(f'[factors] saved to {SAVE_PATH}')


if __name__ == '__main__':
    main()
