'''
Multiple-testing corrections and the deflated Sharpe ratio.

With ~100 runs, some p-values below 5 % are expected by chance.  This file
adjusts the p-values of compare_all.xlsx (run compare_all.py first) and
deflates every run's Sharpe ratio for the number of runs tried.

Sheets of multitest.xlsx (in _output/_comparison/):
  summary   per family of tests: number of tests, and how many are below
            ALPHA raw, after Holm (family-wise error rate) and after
            Benjamini-Hochberg (false discovery rate); split into runs better
            and worse than the comparison
  adjusted  one row per run: for every family the raw p, p_holm and p_bh
            (families: p_bench and p_base of sheet overall, every p_<metric>
            of sheet vs_base)
  dsr       per run: annualised Sharpe ratio of daily excess returns,
            skewness, kurtosis, days, probabilistic Sharpe ratio PSR (the
            probability that the true Sharpe ratio is above 0) and deflated
            Sharpe ratio DSR (the same against the best Sharpe ratio expected
            from N_TRIALS runs of pure noise; Bailey and Lopez de Prado
            2014).  DSR above 0.95: the Sharpe ratio survives the number of
            runs tried.

Usage: set the variables below and run the file.
'''

import numpy as np
import pandas as pd
from scipy import stats

import compare_all as ca
from risk_free_rate import rf_daily

# ----
# Variables
# ----

SOURCE   = ca.SAVE_DIR / 'compare_all.xlsx'
ALPHA    = 0.05
N_TRIALS = None          # number of strategies tried for the DSR; None = number of distinct runs
KEYS     = ['model', 'frequency', 'detector', 'series', 'crisis_target', 'run']

SAVE_PATH = ca.SAVE_DIR / 'multitest.xlsx'


# ----
# Adjustments
# ----

def holm(p: np.ndarray) -> np.ndarray:
    '''Holm (1979) step-down adjusted p-values; NaN stays NaN.'''
    p = np.asarray(p, float)
    out = np.full(len(p), np.nan)
    ok = np.flatnonzero(~np.isnan(p))
    order = ok[np.argsort(p[ok])]
    m = len(order)
    adj = np.maximum.accumulate([(m - i) * p[j] for i, j in enumerate(order)]) if m else []
    out[order] = np.minimum(adj, 1.0)
    return out


def bh(p: np.ndarray) -> np.ndarray:
    '''Benjamini-Hochberg (1995) adjusted p-values (q-values); NaN stays NaN.'''
    p = np.asarray(p, float)
    out = np.full(len(p), np.nan)
    ok = np.flatnonzero(~np.isnan(p))
    order = ok[np.argsort(p[ok])]
    m = len(order)
    if m:
        adj = np.array([p[j] * m / (i + 1) for i, j in enumerate(order)])
        out[order] = np.minimum(np.minimum.accumulate(adj[::-1])[::-1], 1.0)
    return out


def families() -> list[tuple[str, pd.DataFrame, str, str, str]]:
    '''(family name, table, p column, difference column, column tag) of every test family.'''
    ov = pd.read_excel(SOURCE, sheet_name='overall')
    ov = ov.drop_duplicates(subset=['model', 'run', 'frequency'])
    out = []
    bench = ov[~ov['model'].isin([ca.BENCH_MODEL, 'sp500'])]
    out.append(('Sharpe vs benchmark (p_bench)', bench, 'p_bench', 'd_sharpe_bench', 'bench sharpe'))
    try:
        vb = pd.read_excel(SOURCE, sheet_name='vs_base')
    except ValueError:                     # no vs_base sheet: the Sharpe test of sheet overall
        vb = None
        if 'p_base' in ov.columns:
            out.append(('sharpe vs baseline (p_base)', ov[ov['p_base'].notna()], 'p_base',
                        'd_sharpe_base', 'base sharpe'))
    if vb is not None:                     # its p_sharpe is the same test as p_base
        for c in vb.columns:
            if c.startswith('p_') and 'd_' + c[2:] in vb.columns:
                out.append((f'{c[2:]} vs baseline (vs_base)', vb, c, 'd_' + c[2:], 'base ' + c[2:]))
    return out


# better direction of each difference, to split significant results into better / worse
BETTER = {'sharpe': +1, 'sharpe_bench': +1, 'sharpe_base': +1, 'ann_return': +1, 'ann_vol': -1,
          'cum_return': +1, 'max_dd': +1, 'dd_days': -1, 'calmar': +1, 'ulcer': -1,
          'var_95': +1, 'cvar_95': +1, 'beta': -1}


def adjust() -> tuple[pd.DataFrame, pd.DataFrame]:
    summary, wide = [], None
    for name, df, pcol, dcol, tag in families():
        d = df[KEYS + [dcol, pcol]].copy()
        d['p_holm'] = holm(d[pcol].to_numpy())
        d['p_bh'] = bh(d[pcol].to_numpy())
        sign = BETTER.get(dcol[2:], +1) * np.sign(d[dcol])
        row = {'family': name, 'tests': int(d[pcol].notna().sum())}
        for label, col in (('raw', pcol), ('holm', 'p_holm'), ('bh', 'p_bh')):
            sig = d[col] < ALPHA
            row[f'sig_{label}'] = int(sig.sum())
            row[f'better_{label}'] = int((sig & (sign > 0)).sum())
            row[f'worse_{label}'] = int((sig & (sign < 0)).sum())
        if pcol != 'p_sharpe' and 'vs_base' in name:          # bootstrap p-values: floor 1 / (reps + 1)
            floor = 1.0 / (ca.BOOT_REPS + 1)
            row['p_floor'] = floor
            tests = row['tests']
            if floor * tests > ALPHA:
                need = int(np.ceil(tests / ALPHA))
                print(f'[multitest] note: {name}: the smallest bootstrap p is {floor:.4f}, so Holm '
                      f'cannot reject any of {tests} tests; set compare_all.BOOT_REPS to at '
                      f'least {need} (e.g. 5000) to make Holm usable')
        summary.append(row)
        d = d.rename(columns={dcol: f'{tag}: d', pcol: f'{tag}: p', 'p_holm': f'{tag}: p_holm',
                              'p_bh': f'{tag}: p_bh'})
        wide = d if wide is None else wide.merge(d, on=KEYS, how='outer')
    return pd.DataFrame(summary), wide


# ----
# Deflated Sharpe ratio
# ----

def dsr_table() -> pd.DataFrame:
    runs = ca.find_runs()
    rows, seen = [], set()
    for r in sorted(runs, key=ca._sort_key):
        if id(r['returns']) in seen:
            continue
        seen.add(id(r['returns']))
        x = r['returns'] - rf_daily(r['returns'].index).to_numpy()
        T = len(x)
        sr = x.mean() / x.std(ddof=1)
        rows.append({**{k: r[k] for k in KEYS}, 'days': T, 'sr_daily': sr,
                     'sharpe_ann': sr * np.sqrt(ca.TRADING_DAYS_PER_YEAR),
                     'skew': float(stats.skew(x)), 'kurtosis': float(stats.kurtosis(x, fisher=False))})
    df = pd.DataFrame(rows)
    n = N_TRIALS or len(df)
    var_sr = df['sr_daily'].var(ddof=1)
    g = np.euler_gamma
    sr0 = np.sqrt(var_sr) * ((1 - g) * stats.norm.ppf(1 - 1 / n) + g * stats.norm.ppf(1 - 1 / (n * np.e)))

    def psr(row, bench):
        s, k3, k4, T = row['sr_daily'], row['skew'], row['kurtosis'], row['days']
        den = np.sqrt(1 - k3 * s + (k4 - 1) / 4 * s ** 2)
        return float(stats.norm.cdf((s - bench) * np.sqrt(T - 1) / den))

    df['psr'] = df.apply(lambda r: psr(r, 0.0), axis=1)
    df['n_trials'] = n
    df['sr0_ann'] = sr0 * np.sqrt(ca.TRADING_DAYS_PER_YEAR)
    df['dsr'] = df.apply(lambda r: psr(r, sr0), axis=1)
    return df.drop(columns='sr_daily').sort_values('dsr', ascending=False)


# ----
# Main
# ----

def _write(sheets: dict) -> None:
    from openpyxl.utils import get_column_letter
    with pd.ExcelWriter(SAVE_PATH, engine='openpyxl') as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            ws.freeze_panes = 'B2' if name == 'summary' else 'G2'
            for j, col in enumerate(df.columns, 1):
                letter = get_column_letter(j)
                ws.column_dimensions[letter].width = min(max(10, len(str(col)) + 2), 40)
                if df[col].dtype.kind == 'f':
                    for cell in ws[letter][1:]:
                        cell.number_format = '0.0000'


def main() -> None:
    summary, wide = adjust()
    dsr = dsr_table()
    _write({'summary': summary, 'adjusted': wide, 'dsr': dsr})
    cols = ['family', 'tests', 'sig_raw', 'sig_holm', 'sig_bh']
    for fam, tests, raw, ho, b in summary[cols].itertuples(index=False):
        print(f'[multitest] {fam:38s} {tests:3d} tests: {raw:3d} raw, {ho:3d} Holm, {b:3d} BH below {ALPHA}')
    n_ok, n, sr0 = int((dsr['dsr'] > 0.95).sum()), dsr['n_trials'].iloc[0], dsr['sr0_ann'].iloc[0]
    print(f'[multitest] DSR > 0.95: {n_ok} of {len(dsr)} runs (N = {n}, expected max Sharpe of noise {sr0:.2f})')
    print(f'[multitest] saved to {SAVE_PATH}')


if __name__ == '__main__':
    main()
