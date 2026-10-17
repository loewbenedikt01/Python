'''
How accurate are the regime detectors, independent of the portfolio models?

The detector state does not depend on the model (it reads the market), so
every detector is evaluated once, not per run.  Detectors (main series):
changepoint / VIX, hmm / GSPC, wasserstein / GSPC + VIX.  Each uses a
different input series, so differences reflect method plus data.

Two layers
  daily     the detector's own daily state, as stored in
            _regimes/<detector>/regimes_final/<detector>_<series>.csv and read
            by regime_def (same smoothing, SMOOTH_HALFLIFE, and threshold,
            p_crisis > CRISIS_THRESHOLD): crisis on day t if the signal of day
            t-1 is above the threshold.  This is the state a portfolio
            rebalanced every day would hold, with the same one-day convention
            as the applied layer (the signal of date d uses data through d-1;
            the holding decided on d earns from the next day on).
            Frequency column: 'daily'.  3 rows per table.
  applied   the state the portfolios actually hold: the crisis flag of
            regime.csv, set only on the rebalancing dates of the frequency
            (month, quarter or year start) and effective from the next
            trading day, exactly as crisis_analysis_defined.defensive().
            9 rows per table (3 detectors x monthly / quarterly / yearly).
No look-ahead: only stored states are read, nothing is refitted.  Calendar:
the trading days of the portfolio runs, 1998-01-02 .. 2025-12-31; crisis
dates from crises.py (C1 ... C6).  C6 is censored at the end of the data.

Output: _output/_comparison/detector_evaluation.xlsx
  per_crisis   one row per layer x detector x frequency x crisis, all in
               trading days, windows (start, end] as in the crisis scripts:
    exit_lag                    S&P peak -> first defensive day; negative =
                                already defensive at the peak; empty = not
                                defensive peak -> trough
    days_defensive_before_peak  defensive at the peak: days since the switch; else 0
    reentry_lag                 S&P trough -> first day back to normal, only
                                if defensive during the crisis (negative = back
                                before the trough; empty if never back)
    defensive_share_ptt / _ttp / _full
                                share of days defensive, peak -> trough,
                                trough -> break-even, peak -> break-even
    n_switches                  state changes, peak -> break-even
    missed                      never defensive peak -> trough
    late_defensive              not defensive peak -> trough, but defensive
                                trough -> break-even
    defensive_after_trough_days late_defensive only: trough -> first defensive day
    still_defensive_at_end      defensive without a break from its switch in
                                this crisis to the end of the recovery horizon
                                (next crisis peak; C6: end of the data)
    impl_lag_exit, impl_lag_reentry
                                applied rows: applied minus daily exit_lag /
                                reentry_lag, the cost of the rebalancing frequency
    The applied rows are checked against the timing sheets of
    crisis_analysis_all_<frequency>.xlsx (assert).
  full_period  one row per layer x detector x frequency, whole sample:
    total_defensive_share, n_switches, n_episodes, mean_episode_days,
    median_episode_days; n_false_alarm_episodes, false_alarm_days,
    false_alarm_share_of_defensive (false alarm = defensive episode
    overlapping no crisis window from S&P peak - FALSE_ALARM_BUFFER trading
    days to the S&P break-even); n_missed_crises (of 6); day-level
    classification against a crisis label, precision, recall, f1, accuracy,
    false_positive_rate, balanced_accuracy, with the label = S&P peak ->
    trough (suffix _ptt, main) or S&P peak -> break-even (_full, robustness).
  episodes     one row per defensive episode: start, end, days, the crises
               it overlaps (with the buffer) or false_alarm, the S&P 500
               return over its days, ongoing_at_end.
  notes        these definitions.

Usage: run crisis_analysis_defined.py first (for the timing check), then this file.
'''

import os
import sys
from pathlib import Path

import numpy as np
import pandas as pd

import compare_all as ca
import crisis_analysis_defined as cad
from crises import overall, main_crises

sys.path.append(str(Path(__file__).resolve().parents[1] / '_regimes'))
import regime_def


# ----
# Variables
# ----

DETECTORS = [('changepoint', 'vix'), ('hmm', 'gspc'), ('wasserstein', 'gspc_vix')]
FALSE_ALARM_BUFFER = 63          # trading days before the S&P peak that still count as the crisis
SP500_MODEL = 'sp500'            # S&P 500 index run, for the return during an episode

SAVE_PATH = cad.SAVE_DIR / 'detector_evaluation.xlsx'

KEYS = ['layer', 'detector', 'series', 'frequency']
CRISIS_COLS = ['exit_lag', 'days_defensive_before_peak', 'reentry_lag',
               'defensive_share_ptt', 'defensive_share_ttp', 'defensive_share_full',
               'n_switches', 'missed', 'late_defensive', 'defensive_after_trough_days',
               'still_defensive_at_end']
# columns of the timing sheet (crisis_analysis_defined) and their name here
TIMING_MAP = {c: c for c in cad.TIMING_COLS} | {'defensive_share': 'defensive_share_full'}

PCT_COLS = ['defensive_share_ptt', 'defensive_share_ttp', 'defensive_share_full',
            'total_defensive_share', 'false_alarm_share_of_defensive', 'sp500_return']
PCT_COLS += [f'{m}_{s}' for m in ('precision', 'recall', 'f1', 'accuracy',
                                  'false_positive_rate', 'balanced_accuracy') for s in ('ptt', 'full')]


# ----
# States
# ----

def daily_state(detector: str, series: str, calendar: pd.DatetimeIndex) -> pd.Series:
    '''Daily layer: crisis on t if the stored signal of t-1 is above the threshold.'''
    sig = regime_def._signal(detector, series, regime_def.SMOOTH_HALFLIFE)['p_crisis']
    on = sig.reindex(calendar, method='ffill') > regime_def.CRISIS_THRESHOLD   # NaN -> False
    return on.shift(1, fill_value=False).astype(bool)


def applied_states(runs: list[dict]) -> dict:
    '''(detector, series, frequency) -> applied state of one regime run (the same for all).'''
    out = {}
    for r in runs:
        key = (r['detector'], r['series'], r['frequency'])
        if key[:2] in DETECTORS and key not in out:
            out[key] = cad.defensive(r)
    return out


# ----
# Per crisis
# ----

def _pos(idx: pd.DatetimeIndex, date) -> int:
    '''Last position <= date.'''
    return int(idx.searchsorted(pd.Timestamp(date), side='right')) - 1


def _horizon_end(k: int):
    '''Recovery horizon of main crisis k: the next crisis peak (None = end of the data).'''
    return main_crises[k + 1]['peak'] if k + 1 < len(main_crises) else None


def crisis_metrics(d: pd.Series, k: int) -> dict:
    '''CRISIS_COLS for the defensive state d in main crisis k (same rules as the timing sheet).'''
    c = main_crises[k]
    on = d.to_numpy()
    i_p, i_t, i_e = _pos(d.index, c['peak']), _pos(d.index, c['trough']), _pos(d.index, c['even'])
    h = _horizon_end(k)
    i_h = _pos(d.index, h) if h is not None else len(on) - 1

    if on[i_p]:
        j = i_p
        while j > 0 and on[j - 1]:
            j -= 1
        exit_lag, before = j - i_p, i_p - j
    elif on[i_p + 1:i_t + 1].any():
        exit_lag, before = int(np.argmax(on[i_p + 1:i_t + 1])) + 1, 0
    else:
        exit_lag, before = np.nan, 0

    if on[i_t]:
        rest = ~on[i_t + 1:]
        reentry_lag = int(np.argmax(rest)) + 1 if rest.any() else np.nan
    elif on[i_p:i_t + 1].any():
        reentry_lag = i_p + int(np.flatnonzero(on[i_p:i_t + 1])[-1]) + 1 - i_t
    else:
        reentry_lag = np.nan

    ptt, ttp, full = on[i_p + 1:i_t + 1], on[i_t + 1:i_e + 1], on[i_p + 1:i_e + 1]
    missed = bool(not on[i_p:i_t + 1].any())
    late = bool(missed and ttp.any())
    hits = np.flatnonzero(on[i_p:i_e + 1])
    first = i_p + int(hits[0]) if len(hits) else None
    return {
        'exit_lag':                    exit_lag,
        'days_defensive_before_peak':  before,
        'reentry_lag':                 reentry_lag,
        'defensive_share_ptt':         float(ptt.mean()) if len(ptt) else np.nan,
        'defensive_share_ttp':         float(ttp.mean()) if len(ttp) else np.nan,
        'defensive_share_full':        float(full.mean()) if len(full) else np.nan,
        'n_switches':                  int(np.abs(np.diff(on[i_p:i_e + 1].astype(int))).sum()),
        'missed':                      missed,
        'late_defensive':              late,
        'defensive_after_trough_days': int(np.argmax(ttp)) + 1 if late else np.nan,
        'still_defensive_at_end':      bool(first is not None and on[first:i_h + 1].all()),
    }


# ----
# Full period
# ----

def episodes(d: pd.Series) -> list[tuple[int, int]]:
    '''(first, last) positions of every run of defensive days.'''
    on = np.r_[False, d.to_numpy(), False].astype(int)
    starts = np.flatnonzero(np.diff(on) == 1)
    ends = np.flatnonzero(np.diff(on) == -1) - 1
    return list(zip(starts, ends))


def crisis_windows(idx: pd.DatetimeIndex) -> list[tuple[str, int, int]]:
    '''(label, first, last) positions of each crisis window incl. the pre-peak buffer.'''
    return [(f'C{k}', max(0, _pos(idx, c['peak']) - FALSE_ALARM_BUFFER), _pos(idx, c['even']))
            for k, c in enumerate(main_crises, 1)]


def label(idx: pd.DatetimeIndex, end: str) -> np.ndarray:
    '''Crisis label per day: (S&P peak, S&P trough] (end='trough') or (peak, break-even] ('even').'''
    y = np.zeros(len(idx), dtype=bool)
    for c in main_crises:
        y[_pos(idx, c['peak']) + 1:_pos(idx, c[end]) + 1] = True
    return y


def classification(on: np.ndarray, y: np.ndarray, suffix: str) -> dict:
    tp, fp = int((on & y).sum()), int((on & ~y).sum())
    fn, tn = int((~on & y).sum()), int((~on & ~y).sum())
    div = lambda a, b: a / b if b else np.nan
    prec, rec, tnr = div(tp, tp + fp), div(tp, tp + fn), div(tn, tn + fp)
    return {
        f'precision_{suffix}':           prec,
        f'recall_{suffix}':              rec,
        f'f1_{suffix}':                  div(2 * prec * rec, prec + rec) if tp else 0.0,
        f'accuracy_{suffix}':            div(tp + tn, len(on)),
        f'false_positive_rate_{suffix}': div(fp, fp + tn),
        f'balanced_accuracy_{suffix}':   (rec + tnr) / 2,
    }


def full_period(d: pd.Series, per_crisis: list[dict], sp: pd.Series, keys: dict):
    '''(full_period row, episode rows) for one state d.'''
    on = d.to_numpy()
    wins = crisis_windows(d.index)
    eps = episodes(d)
    ep_rows = []
    for s, e in eps:
        hit = [lab for lab, lo, hi in wins if s <= hi and e >= lo]
        ep_rows.append({**keys, 'start': d.index[s], 'end': d.index[e], 'days': e - s + 1,
                        'overlap': ', '.join(hit) if hit else 'false_alarm',
                        'sp500_return': float(np.expm1(sp.reindex(d.index[s:e + 1]).sum())),
                        'ongoing_at_end': bool(e == len(on) - 1)})
    lengths = np.array([e - s + 1 for s, e in eps])
    fa = np.array([r['overlap'] == 'false_alarm' for r in ep_rows], dtype=bool)
    row = {
        **keys,
        'total_defensive_share':          float(on.mean()),
        'n_switches':                     int(np.abs(np.diff(on.astype(int))).sum()),
        'n_episodes':                     len(eps),
        'mean_episode_days':              float(lengths.mean()) if len(eps) else np.nan,
        'median_episode_days':            float(np.median(lengths)) if len(eps) else np.nan,
        'n_false_alarm_episodes':         int(fa.sum()),
        'false_alarm_days':               int(lengths[fa].sum()) if len(eps) else 0,
        'false_alarm_share_of_defensive': float(lengths[fa].sum() / on.sum()) if on.sum() else np.nan,
        'n_missed_crises':                int(sum(m['missed'] for m in per_crisis)),
    }
    row.update(classification(on, label(d.index, 'trough'), 'ptt'))
    row.update(classification(on, label(d.index, 'even'), 'full'))
    return row, ep_rows


# ----
# Check against the timing sheets
# ----

def check_timing(per_crisis: pd.DataFrame) -> int:
    '''Assert that the applied rows equal the timing sheets; returns the number of rows checked.'''
    n = 0
    for freq in cad.FREQ_ORDER:
        path = cad.SAVE_DIR / f'crisis_analysis_all_{freq}.xlsx'
        timing = pd.read_excel(path, sheet_name='timing')
        mine = per_crisis[(per_crisis['layer'] == 'applied') & (per_crisis['frequency'] == freq)]
        merged = timing.merge(mine, on=['frequency', 'detector', 'series', 'crisis'],
                              suffixes=('_t', ''), validate='one_to_one')
        assert len(merged) == len(timing) == len(mine), f'{freq}: rows do not line up'
        for t_col, col in TIMING_MAP.items():
            if t_col not in timing:
                continue
            a = merged[f'{t_col}_t' if t_col == col else t_col].astype(float).to_numpy()
            b = merged[col].astype(float).to_numpy()
            assert np.allclose(a, b, equal_nan=True), f'{freq} {col}: differs from the timing sheet'
        n += len(merged)
    return n


# ----
# Output
# ----

NOTES = [
    'Detector evaluation, independent of the portfolio models (the detector reads the market).',
    'Detectors: changepoint = VIX, HMM = S&P 500 (GSPC), Wasserstein k-means = S&P 500 + VIX. '
    'Each uses a different input series, so differences reflect method plus data.',
    'Layer daily: the stored daily signal (regime_def, same smoothing and threshold); crisis on day t '
    'if the signal of t-1 is above the threshold, i.e. a portfolio rebalanced every day.',
    'Layer applied: the state the portfolios hold, set on the rebalancing dates of the frequency and '
    'effective from the next trading day (crisis_analysis_defined.defensive).',
    'implementation lag = applied minus daily exit_lag / reentry_lag: the cost of the rebalancing frequency.',
    'Windows (start, end] in trading days; crisis dates from crises.py (S&P peak, trough, break-even).',
    'exit_lag: S&P peak -> first defensive day (negative = already defensive at the peak; empty = not '
    'defensive peak -> trough). reentry_lag: S&P trough -> first day back to normal (only if defensive '
    'during the crisis; empty if never back).',
    'missed: never defensive peak -> trough. late_defensive: missed, but defensive trough -> break-even. '
    'still_defensive_at_end: no break from the switch to the next crisis peak (C6: end of the data).',
    f'False alarm: a defensive episode overlapping no window from S&P peak - {FALSE_ALARM_BUFFER} '
    'trading days to the S&P break-even.',
    'Classification per day against the label S&P peak -> trough (_ptt, main) and peak -> break-even '
    '(_full, robustness).',
    'No look-ahead: only stored detector states are read; nothing is refitted.',
    'C6 is censored at the end of the data (2025-12-31).',
]


def write(sheets: dict, path: Path) -> None:
    '''Write to a temp file, then rename, so a locked target is never half-written.'''
    from openpyxl.utils import get_column_letter
    tmp = path.with_name(path.stem + '.tmp.xlsx')
    with pd.ExcelWriter(tmp, engine='openpyxl') as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            n_keys = sum(c in KEYS + ['crisis'] for c in df.columns)
            if n_keys:
                ws.freeze_panes = get_column_letter(n_keys + 1) + '2'
            for j, col in enumerate(df.columns, 1):
                letter = get_column_letter(j)
                ws.column_dimensions[letter].width = 110 if name == 'notes' else min(max(10, len(str(col)) + 2), 32)
                fmt = ('0.00%' if col in PCT_COLS else 'yyyy-mm-dd' if col in ('start', 'end')
                       else '0.000' if df[col].dtype.kind == 'f' and 'lag' not in col and 'days' not in col
                       else None)
                if fmt:
                    for cell in ws[letter][1:]:
                        cell.number_format = fmt
    try:
        os.replace(tmp, path)
    except PermissionError:
        print(f'[detector_evaluation] {path.name} is open (locked): close it and run again; '
              f'the new results are in {tmp.name}')
        return
    print(f'[detector_evaluation] saved to {path}')


def main() -> None:
    runs = ca.find_runs()
    bench = next(r for r in runs if r['model'] == cad.BENCHMARK and r['frequency'] == 'monthly')
    w = overall[0]
    calendar = ca._slice(bench['returns'], w['peak'], w['even']).index
    sp = next(r for r in runs if r['model'] == SP500_MODEL)['returns']
    applied = applied_states(runs)

    states = []                                  # (keys, state)
    for det, ser in DETECTORS:
        states.append(({'layer': 'daily', 'detector': det, 'series': ser, 'frequency': 'daily'},
                       daily_state(det, ser, calendar)))
    for freq in cad.FREQ_ORDER:
        for det, ser in DETECTORS:
            d = applied.get((det, ser, freq))
            if d is None:
                print(f'[detector_evaluation] no regime run for {det} {ser} {freq}: skipped')
                continue
            states.append(({'layer': 'applied', 'detector': det, 'series': ser, 'frequency': freq},
                           d.reindex(calendar, fill_value=False)))

    crisis_rows, full_rows, ep_rows = [], [], []
    for keys, d in states:
        pc = [crisis_metrics(d, k) for k in range(len(main_crises))]
        crisis_rows += [{**keys, 'crisis': f'C{k + 1}', **m} for k, m in enumerate(pc)]
        row, eps = full_period(d, pc, sp, keys)
        full_rows.append(row)
        ep_rows += eps

    per_crisis = pd.DataFrame(crisis_rows)
    daily = per_crisis[per_crisis['layer'] == 'daily'].set_index(['detector', 'crisis'])
    is_applied = per_crisis['layer'] == 'applied'
    for lag, col in (('exit_lag', 'impl_lag_exit'), ('reentry_lag', 'impl_lag_reentry')):
        ref = per_crisis.set_index(['detector', 'crisis']).index.map(daily[lag])
        per_crisis[col] = (per_crisis[lag] - np.asarray(ref, dtype=float)).where(is_applied)

    n = check_timing(per_crisis)
    print(f'[detector_evaluation] applied layer matches the timing sheets ({n} rows)')

    write({'per_crisis': per_crisis, 'full_period': pd.DataFrame(full_rows),
           'episodes': pd.DataFrame(ep_rows), 'notes': pd.DataFrame({'note': NOTES})}, SAVE_PATH)


if __name__ == '__main__':
    main()
