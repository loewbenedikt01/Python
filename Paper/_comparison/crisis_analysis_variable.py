'''
Crisis performance of the baseline portfolios with portfolio-specific windows.

Same portfolios, files and sheets as crisis_analysis_defined.py (baseline runs
= no changepoint / hmm / wasserstein in the name; one workbook per frequency;
sheets overall, C1 ... C7 from crises.main_crises), but only the start of each
crisis is taken from crises.py.  From there every portfolio gets its own
episode:
  start     the crisis peak date of crises.py; portfolio value = 1 (anchor)
  trough    the portfolio's lowest value in (start, search end], search end =
            defined recovery date + TROUGH_BUFFER_MONTHS (at most the next
            crisis peak).  A trough on the last day of the search window is
            flagged (trough_at_boundary): read its depth as "at least".
  recovery  the first day after the trough with value >= 1 (back at the
            anchor), searched up to the next crisis peak (last crisis: end of
            the data); none = not recovered
  own window  start -> recovery (-> recovery search end if not recovered)
A portfolio whose lowest value stays above 1 - NO_DD_TOL has no drawdown: no
trough or recovery date, ptt_days = ttp_days = 0, own window = defined window.
Not recovered = censored: full_days, full_days_underwater and ttp_days are then
counted up to the recovery search end and mean "at least this many days".
C1 ... C6 are censored by the next crisis; C7 by the end of the data
(2025-12-31), so an unrecovered C7 portfolio may simply not have had the time.
Sheet overall: anchoring at the 1998 start would only catch an early dip, so
the start is the portfolio's own high-water mark before its deepest drawdown
of the full period, searched to the end of the data.  These episodes differ
between portfolios (mcap: dot-com peak 2000 -> GFC trough 2009, most others:
GFC only), so ptt_days_vs_defined and the benchmark columns stay empty there;
the overall sheet carries the hit rates over C1 ... C7 instead.

Output (as in crisis_analysis_defined.py, see there for roles and fallbacks)
  1. baselines: _output/_comparison/crisis_analysis_variable_<frequency>.xlsx
  2. per regime model and frequency (benchmark, baseline, fallbacks, regime runs):
     _output/_comparison/crisis_analysis_regime/crisis_analysis_variable_<model>_<frequency>.xlsx
  3. all runs, one workbook per frequency:
     _output/_comparison/crisis_analysis_variable_all_<frequency>.xlsx
  2. and 3. carry all columns of 1., plus on the crisis sheets
       full_d_base_ulcer    full_ulcer_common - base_run (negative = less pain)
       ptt_d_base_depth     ptt_depth - base_run (positive = shallower)
       ttp_d_base_days      ttp_days - base_run (negative = faster); empty if
                            either has not recovered (ttp_d_base_censored)
                            (0 for the baseline, empty for the benchmark)
       full_d_fb_ulcer, ptt_d_fb_depth, ttp_d_fb_days, ttp_d_fb_censored
                            the same against fallback_run (the crisis target
                            held all the time)
       missed_rally, pre_peak_d_base
                            cost of being defensive before the peak
                            (crisis_analysis_defined.add_switch_cost)
     on the overall sheet hit_rate_base_* / hit_count_base_* and
     hit_rate_fb_* / hit_count_fb_*: hit rates against base_run / fallback_run,
     and the sheets timing (detector timing, once per detector, series,
     frequency and crisis; switches only on rebalancing dates) and, in 3.,
     detector_avg (deltas averaged over the models per detector).
  Every number depends only on the run and the crisis, not on the workbook.
Rows: one per portfolio (model, run, frequency).  Days = trading days.

Columns (bench = BENCHMARK run of the same frequency)
  peak_date               start (the defined crisis peak; overall: own peak)
  trough_date, recovery_date   the portfolio's own dates
  full_days               length of the own window
  full_anchored_dd        trough value / anchor - 1
  full_hwm_dd             max drawdown from the running high-water mark
                          (can be deeper than anchored if the portfolio rose first)
  full_days_underwater    days below the anchor (start value) in the own window
  full_ulcer              Ulcer index over the own window, drawdowns from the
                          running high (standard definition, metrics.ulcer_index);
                          own windows differ in length, so a longer recovery
                          can lower it: not used for comparisons
  full_ulcer_common       Ulcer index over a fixed horizon per crisis:
                          defined start -> full_common_end = the next crisis
                          peak (C7: end of the data), the same as the
                          recovery search end; drawdowns from the running
                          high until the own recovery, 0 after.  Same number
                          of days for every portfolio, so a longer recovery
                          raises it, and one value per run and crisis in
                          every workbook.  Crisis sheets only.
  full_common_end         end of that common horizon
  full_no_drawdown        True if the lowest value stays above 1 - NO_DD_TOL
  full_not_recovered      True if not back at the anchor by the recovery
                          search end (day counts censored)
  full_gap_at_cutoff      not recovered only: value at the recovery search end
                          / anchor - 1 (how far below the start it still was)
  ptt_days                start -> own trough
  ptt_days_vs_defined     ptt_days - days from start to the defined trough
                          (negative = bottomed earlier)
  ptt_depth               anchored drawdown at the trough (= full_anchored_dd)
  ptt_max_dd              max drawdown (high-water mark) start -> trough
  ptt_trough_at_boundary  True if the trough is the last day of the search window
  ttp_days                own trough -> own recovery (censored if not recovered)
  ttp_days_vs_bench       ttp_days - bench ttp_days (negative = recovered
                          faster); empty if either has not recovered
  ttp_vs_bench_censored   True if ttp_days_vs_bench is empty for that reason
  ttp_return_to_even      return from the own trough back to the anchor,
                          1 / trough value - 1
  hit_rate_full, hit_rate_ptt, hit_rate_ttp
                          overall sheet only: share of the crises C1 ... C7 in
                          which the portfolio beats the benchmark (empty for
                          the benchmark).  full: lower full_ulcer_common; ptt:
                          shallower ptt_depth; ttp: fewer ttp_days (recovered
                          beats not recovered; both not recovered = crisis
                          left out of the share)
  hit_count_full, hit_count_ptt, hit_count_ttp
                          the same as hits / comparable crises, e.g. 2/5

Usage: run the file.
'''

import numpy as np
import pandas as pd

import compare_all as ca
from crises import overall, main_crises
from metrics import maximum_drawdown, ulcer_index
from crisis_analysis_defined import KEY_COLS, REFS, _base_value, add_hit_rates, run_all


# ----
# Variables
# ----

TROUGH_BUFFER_MONTHS = 6       # trough search past the defined recovery date
NO_DD_TOL            = 0.01    # lowest value above 1 - NO_DD_TOL = no drawdown

PCT_COLS  = ['full_anchored_dd', 'full_hwm_dd', 'full_ulcer', 'full_ulcer_common', 'full_gap_at_cutoff',
             'ptt_depth', 'ptt_max_dd',
             'ttp_return_to_even', 'hit_rate_full', 'hit_rate_ptt', 'hit_rate_ttp',
             'hit_rate_base_full', 'hit_rate_base_ptt', 'hit_rate_base_ttp',
             'hit_rate_fb_full', 'hit_rate_fb_ptt', 'hit_rate_fb_ttp',
             'full_d_base_ulcer', 'ptt_d_base_depth', 'full_d_fb_ulcer', 'ptt_d_fb_depth',
             'missed_rally', 'pre_peak_d_base', 'defensive_share']
DATE_COLS = ['peak_date', 'trough_date', 'recovery_date', 'full_common_end']


# ----
# Episode
# ----

def _with_anchor(lvl: pd.Series) -> pd.Series:
    '''Level series with the anchor value 1 prepended (for drawdowns from the start).'''
    return pd.concat([pd.Series([1.0]), lvl.reset_index(drop=True)], ignore_index=True)


def episode(returns: pd.Series, c: dict, rec_end=None, own_peak: bool = False) -> dict:
    '''
    The portfolio's own crisis episode from the defined start c['peak'].
    rec_end: last day of the recovery search (the next crisis peak; None = end
    of the data).  own_peak (overall sheet): start at the high-water mark
    before the portfolio's deepest drawdown in the window instead.
    '''
    r = returns[returns.index > pd.Timestamp(c['peak'])]
    lvl = np.exp(r.cumsum())
    if rec_end is not None:
        lvl = lvl[lvl.index <= pd.Timestamp(rec_end)]
    peak_date = pd.Timestamp(c['peak'])
    even = pd.Timestamp(c['even'])
    if own_peak:
        win = lvl[lvl.index <= even]
        t = (win / win.cummax()).idxmin()
        peak_date = win.loc[:t].idxmax()
        lvl = lvl[lvl.index > peak_date] / lvl.loc[peak_date]
        trough_end = even
    else:
        trough_end = even + pd.DateOffset(months=TROUGH_BUFFER_MONTHS)
    win = lvl[lvl.index <= trough_end]
    if len(win) < 2:
        return {}
    defined_ptt = np.nan if own_peak else int((lvl.index <= pd.Timestamp(c['trough'])).sum())
    t_date, t_val = win.idxmin(), float(win.min())

    if t_val >= 1.0 - NO_DD_TOL:                      # no (meaningful) drawdown
        own = lvl[lvl.index <= even]
        ptt_days, ttp_days, t_date, rec_date = 0, 0, None, None
        no_dd, not_rec, at_boundary = True, False, False
    else:
        at_boundary = bool(t_date == win.index[-1])
        ptt_days = lvl.index.get_loc(t_date) + 1
        after = lvl[lvl.index > t_date]
        back = after[after >= 1.0]
        no_dd, not_rec = False, back.empty
        rec_date = None if not_rec else back.index[0]
        own = lvl if not_rec else lvl[lvl.index <= rec_date]
        ttp_days = len(own) - ptt_days                # lower bound if not recovered

    own_lvl = _with_anchor(own)
    # drawdowns from the running high over the own window (0 after recovery,
    # see common_ulcer); none for no drawdown
    dd = (own_lvl / own_lvl.cummax() - 1.0).iloc[1:].set_axis(own.index)
    return {
        '_dd':                    dd if not no_dd else pd.Series(dtype=float),
        '_end':                   own.index[-1],
        'peak_date':              peak_date,
        'trough_date':            t_date,
        'recovery_date':          rec_date,
        'full_days':              len(own),
        'full_anchored_dd':       min(t_val - 1.0, 0.0),     # never above the start: 0
        'full_hwm_dd':            maximum_drawdown(own_lvl),
        'full_days_underwater':   int((own < 1.0).sum()),
        'full_ulcer':             ulcer_index(own_lvl),
        'full_no_drawdown':       no_dd,
        'full_not_recovered':     not_rec,
        'full_gap_at_cutoff':     float(own.iloc[-1]) - 1.0 if not_rec else np.nan,
        'ptt_days':               ptt_days,
        'ptt_days_vs_defined':    ptt_days - defined_ptt,
        'ptt_depth':              min(t_val - 1.0, 0.0),
        'ptt_max_dd':             maximum_drawdown(_with_anchor(lvl.iloc[:ptt_days])),
        'ptt_trough_at_boundary': at_boundary,
        'ttp_days':               ttp_days,
        'ttp_return_to_even':     0.0 if no_dd else 1.0 / t_val - 1.0,
    }


def common_ulcer(returns: pd.Series, ep: dict, start, end) -> float:
    '''
    Ulcer index over the common horizon (start, end] of a crisis: the
    portfolio's drawdowns from the running high until its own recovery, 0
    after it, so every portfolio has the same number of days and a longer
    recovery raises the value.
    '''
    days = returns.index[(returns.index > pd.Timestamp(start)) & (returns.index <= end)]
    dd = ep['_dd'].reindex(days).fillna(0.0).to_numpy()
    return float(np.sqrt(np.mean(np.concatenate([[0.0], dd]) ** 2)))


def _faster(r, b, days: str) -> float:
    '''
    Back at the start value sooner than the reference (censored = not
    recovered by the search end).  Recovered vs not recovered decides; both
    not recovered: not comparable (NaN).  No drawdown = nothing to recover
    from (days = 0, not censored): a hit against a reference with a
    drawdown; both without drawdown: not comparable (NaN).
    '''
    if bool(r['full_no_drawdown']) and bool(b['full_no_drawdown']):
        return np.nan
    r_cens, b_cens = bool(r['full_not_recovered']), bool(b['full_not_recovered'])
    if r_cens and b_cens:
        return np.nan
    if r_cens != b_cens:
        return float(b_cens)
    return float(r[days] < b[days])


# Beats the benchmark: full = lower Ulcer index over the common horizon (less
# pain from depth and duration combined), ptt = shallower drop to the own
# trough, ttp = faster recovery from the trough
HIT_TESTS = {
    'full': lambda r, b: float(r['full_ulcer_common'] < b['full_ulcer_common']),
    'ptt':  lambda r, b: float(r['ptt_depth'] > b['ptt_depth']),
    'ttp':  lambda r, b: _faster(r, b, 'ttp_days'),
}


def vs_bench(row: dict, b: dict) -> None:
    '''ttp_days_vs_bench, empty (and flagged) if either has not recovered.'''
    if 'ttp_days' not in row or 'ttp_days' not in b:
        return
    censored = row['full_not_recovered'] or b['full_not_recovered']
    row['ttp_days_vs_bench'] = np.nan if censored else row['ttp_days'] - b['ttp_days']
    row['ttp_vs_bench_censored'] = censored


# ----
# Main
# ----

def add_base_deltas(sheets: dict) -> None:
    '''
    Crisis sheets, for each reference of REFS (base_run -> _d_base_,
    fallback_run -> _d_fb_): full_d_*_ulcer, ptt_d_*_depth, ttp_d_*_days (and
    ttp_d_*_censored); 0 for the reference itself, empty without one.
    ttp_d_*_days is empty if either has not recovered.
    '''
    for name, df in sheets.items():
        if name == 'overall' or 'ttp_days' not in df:
            continue
        for ref, suf in REFS:
            if ref not in df:
                continue
            ref_val = lambda col: _base_value(df, col, ref)
            keep = df[ref].notna()
            censored = df['full_not_recovered'].astype(bool) | ref_val('full_not_recovered').eq(True)
            df[f'full_d_{suf}_ulcer'] = df['full_ulcer_common'] - ref_val('full_ulcer_common')
            df[f'ptt_d_{suf}_depth'] = df['ptt_depth'] - ref_val('ptt_depth')
            df[f'ttp_d_{suf}_days'] = (df['ttp_days'] - ref_val('ttp_days')).where(~censored)
            df[f'ttp_d_{suf}_censored'] = censored.astype(object).where(keep)


def _write(sheets: dict, path, n_keys: int = len(KEY_COLS)) -> None:
    from openpyxl.utils import get_column_letter
    with pd.ExcelWriter(path, engine='openpyxl') as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            ws.freeze_panes = get_column_letter(n_keys + 1) + '2'
            for j, col in enumerate(df.columns, 1):
                letter = get_column_letter(j)
                ws.column_dimensions[letter].width = min(max(10, len(str(col)) + 2), 32)
                fmt = ('0.00%' if col in PCT_COLS else 'yyyy-mm-dd' if col in DATE_COLS
                       else '0.000' if df[col].dtype.kind == 'f' and 'days' not in col else None)
                if fmt:
                    for cell in ws[letter][1:]:
                        cell.number_format = fmt


# (sheet, crisis, recovery search end, own peak)
WINDOWS = [('overall', overall[0], None, True)]
for _i, _c in enumerate(main_crises):
    WINDOWS.append((f'C{_i + 1}', _c,
                    main_crises[_i + 1]['peak'] if _i + 1 < len(main_crises) else None, False))


def build_sheets(group: list[dict], bench: dict, keys=KEY_COLS, info=None) -> dict:
    sheets = {}
    for name, c, rec_end, own_peak in WINDOWS:
        eps = [episode(run['returns'], c, rec_end, own_peak) for run in group]
        b = eps[group.index(bench)]
        rows = []
        for run, ep, i in zip(group, eps, info or [{}] * len(group)):
            row = {**{k: i[k] if k in i else run[k] for k in keys}, **ep}
            if ep and not own_peak:
                # fixed horizon: up to the next crisis peak (last crisis: end of the data)
                horizon = pd.Timestamp(rec_end) if rec_end else run['returns'].index[-1]
                row['full_ulcer_common'] = common_ulcer(run['returns'], ep, c['peak'], horizon)
                row['full_common_end'] = horizon
            if not own_peak:
                vs_bench(row, b)
            else:                     # overall: own episodes differ, no comparison
                row.update({'ttp_days_vs_bench': np.nan, 'ttp_vs_bench_censored': None})
            rows.append({k: v for k, v in row.items() if not k.startswith('_')})
        df = pd.DataFrame(rows)
        for col in ('ttp_vs_bench_censored', 'ttp_days_vs_bench'):   # next to ttp_days
            if col in df:
                df.insert(df.columns.get_loc('ttp_days') + 1, col, df.pop(col))
        for col in ('full_common_end', 'full_ulcer_common'):         # next to full_ulcer
            if col in df:
                df.insert(df.columns.get_loc('full_ulcer') + 1, col, df.pop(col))
        sheets[name] = df
    add_hit_rates(sheets, HIT_TESTS)
    return sheets


def main() -> None:
    run_all('crisis_analysis_variable', 'crisis_analysis_variable',
            build_sheets, add_base_deltas, HIT_TESTS, _write)


if __name__ == '__main__':
    main()
