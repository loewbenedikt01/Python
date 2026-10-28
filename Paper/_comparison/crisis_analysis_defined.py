'''
Crisis performance of the portfolios over the defined crisis windows.

1. Baselines: every run without a regime detector (no changepoint / hmm /
wasserstein in the name), e.g. market_cap, equal_weight, mvo, hrp, xgb, rf,
lstm; the S&P 500 index is left out (SKIP_MODELS).  One workbook per
frequency: _output/_comparison/crisis_analysis_<frequency>.xlsx

2. Regime models: for every model of REGIME_MODELS and frequency one workbook
_output/_comparison/crisis_analysis_regime/crisis_analysis_<model>_<frequency>.xlsx
with the benchmark (mcap), the model's baseline, the fallback portfolios of
its crisis targets held all the time (FALLBACKS: always min_var for MVO / HRP,
pure HRP for RF / XGB / LSTM; does switching add anything over holding the
fallback?) and all its regime runs.  Column role = benchmark / baseline /
fallback / regime, base_run = the model's baseline.

3. All runs: one workbook per frequency,
_output/_comparison/crisis_analysis_all_<frequency>.xlsx (sheets overall,
C1 ... C7, timing, detector_avg): every run of the
frequency (benchmark, all baselines, controls, all regime runs), each
compared with the benchmark and with its own baseline (base_run).

2. and 3. carry all columns of 1., plus
  fallback_run          regime runs: the run holding the fallback of their
                        crisis target all the time (none for cash)
  full_d_base, ptt_d_base, ttp_d_base
                        return against base_run over the same window,
                        relative: (1 + r) / (1 + r_base) - 1; 0 for the
                        baseline, empty for the benchmark (overall: full only)
  full_d_fb, ptt_d_fb, ttp_d_fb
                        the same against fallback_run: an overlay that falls
                        back to HRP has to beat pure HRP, not only its base
                        model, before the switching adds value
  missed_rally, pre_peak_d_base
                        crisis sheets, regime runs already defensive at the
                        peak: base model return from the switch to the peak,
                        and the run against its base over that interval
                        (the cost of switching early; add_switch_cost)
  hit_rate_base_*, hit_count_base_*, hit_rate_fb_*, hit_count_fb_*
                        overall sheet: hit rates against base_run /
                        fallback_run instead of the benchmark
plus the sheet timing: detector timing per detector, series, frequency and
crisis (detector_diagnostics: exit_lag, reentry_lag,
days_defensive_before_peak, late_defensive, defensive_after_trough_days,
defensive_share, n_switches).  The detector reads the market, not the
portfolio, so the timing is the same for every model and crisis target and
is reported once (identical checks this).  3. also has detector_avg: the
deltas against base and fallback averaged over the models per detector,
series, crisis target, frequency and crisis.  The models of one detector
share its switching dates and are not independent observations: test these
averages, or cluster by detector x frequency, not the model rows.
Every number depends only on the run and the window, not on the workbook
it appears in, so the three outputs agree.

Methodology note: a detector can only switch on the rebalancing dates of its
frequency (regime.csv holds one decision per rebalancing date; the holding
then applies until the next one).  With yearly rebalancing it reacts at most
once a year, so its timing is coarse by design.  The detector signals are
point-in-time (value at d from data through d-1; HMM: recursive filter,
WK-means: monthly refits on past windows, changepoint: sequential cpm alarms),
so a decision never uses data after its date.

Runs are found as in compare_all.py
(_output/<model>/<run>_<frequency>/raw/daily_returns.csv).

Sheets
  overall      crises.overall: full-window metrics plus CAGR, Sharpe,
               Sortino, Calmar and turnover
  C1 ... C7    crises.main_crises: full window (peak -> recovery), PTT
               (peak -> trough) and TTP (trough -> recovery)
Rows: one per portfolio (model, run, frequency).  Every window is (start, end]
and the portfolio value is 1 at its start.

Excess return = relative: (1 + r) / (1 + r_bench) - 1, the end value against
the benchmark's.  Unlike r - r_bench it stays comparable over long windows
with large compounded returns (e.g. dot-com, peak -> recovery).

Columns (bench = BENCHMARK run of the same frequency, same window)
  full_cum_return       cumulative return
  full_excess_return    relative excess return
  full_max_dd           maximum drawdown
  full_ulcer            Ulcer index
  full_ann_vol          annualised volatility
  full_underwater_share share of days below the running peak
  full_recovered        True if the value at the end >= the value at the
                        crisis start (full window only: in PTT / TTP the value
                        is reset to 1 at the phase start, so >= 1 means nothing)
  full_cagr, full_sharpe, full_sortino, full_calmar, full_turnover_pa
                        overall sheet only (metrics.py; turnover per year
                        without the initial build, as in compare_all.py)
  hit_rate_full, hit_rate_ptt, hit_rate_ttp
                        overall sheet only: share of the crises C1 ... C7 in
                        which the excess return of that phase is > 0
                        (empty for the benchmark)
  hit_count_full, hit_count_ptt, hit_count_ttp
                        the same as hits / crises, e.g. 4/6
  ptt_cum_return        cumulative return peak -> trough
  ptt_excess_return     relative excess return peak -> trough
  ptt_down_capture      ptt_cum_return / bench ptt_cum_return (< 1 = lost less)
  ptt_max_dd            maximum drawdown peak -> trough
  ttp_cum_return        cumulative return trough -> recovery
  ttp_excess_return     relative excess return trough -> recovery
  ttp_recovery_capture  ttp_cum_return / bench ttp_cum_return (> 1 = recovered more)
  ttp_max_dd            maximum drawdown trough -> recovery

Usage: run the file.
'''

import numpy as np
import pandas as pd

import compare_all as ca
from crises import overall, main_crises
from metrics import (
    annualized_return,
    annualized_volatility,
    calmar_ratio,
    cumulative_return,
    maximum_drawdown,
    sharpe_ratio,
    sortino_ratio,
    ulcer_index,
)


# ----
# Variables
# ----

BENCHMARK   = 'market_cap'               # model folder of the benchmark run
SKIP_MODELS = {'sp500'}                  # model folders left out (S&P 500 index)
MODEL_ORDER = ['market_cap', 'equal_weight', 'mvo', 'hrp', 'xgb', 'rf', 'lstm']
FREQ_ORDER  = ['monthly', 'quarterly', 'yearly']
REGIME_MODELS = ['mvo', 'hrp', 'xgb', 'rf', 'lstm']    # one regime workbook per model

SAVE_DIR    = ca.SAVE_DIR
REGIME_DIR  = SAVE_DIR / 'crisis_analysis_regime'

KEY_COLS    = ['model', 'run', 'frequency']
REGIME_KEYS = ['model', 'run', 'frequency', 'role', 'base_run', 'fallback_run',
               'detector', 'series', 'crisis_target']

# Fallback of each crisis target: the portfolio switched into, held all the
# time (does switching add anything over always holding it?)
FALLBACKS = {
    'min_var': lambda r: r['crisis_target'] == 'always min_var',
    'hrp':     lambda r: r['model'] == 'hrp' and r['detector'] == 'none' and r['crisis_target'] == '',
}
# Detector timing per crisis (detector_diagnostics): a property of detector x
# series x frequency, reported once in the timing sheet, not per model row
TIMING_COLS = ['exit_lag', 'reentry_lag', 'days_defensive_before_peak', 'late_defensive',
               'defensive_after_trough_days', 'defensive_share', 'n_switches']
TIMING_KEYS = ['frequency', 'detector', 'series', 'crisis']

PCT_COLS = ['full_cum_return', 'full_excess_return', 'full_max_dd', 'full_ulcer', 'full_ann_vol',
            'full_underwater_share', 'full_cagr', 'full_turnover_pa',
            'hit_rate_full', 'hit_rate_ptt', 'hit_rate_ttp',
            'hit_rate_base_full', 'hit_rate_base_ptt', 'hit_rate_base_ttp',
            'hit_rate_fb_full', 'hit_rate_fb_ptt', 'hit_rate_fb_ttp',
            'ptt_cum_return', 'ptt_excess_return', 'ptt_max_dd',
            'ttp_cum_return', 'ttp_excess_return', 'ttp_max_dd',
            'full_d_base', 'ptt_d_base', 'ttp_d_base', 'full_d_fb', 'ptt_d_fb', 'ttp_d_fb',
            'missed_rally', 'pre_peak_d_base', 'defensive_share']

WINDOWS = ([('overall', overall[0], False)]
           + [(f'C{i}', c, True) for i, c in enumerate(main_crises, 1)])


# ----
# Metrics
# ----

def _ratio(a: float, b: float) -> float:
    return a / b if b != 0 and np.isfinite(b) else np.nan


def _excess(cum: float, b_cum: float) -> float:
    '''Relative excess return: end value against the benchmark's end value.'''
    return (1 + cum) / (1 + b_cum) - 1


def _turnover_pa(run: dict, r: pd.Series, start, end) -> float:
    '''Turnover per year in (start, end], without the initial build (as compare_all).'''
    to = run['turnover']
    if to is None or len(to) < 2:
        return np.nan
    years = (r.index[-1] - r.index[0]).days / 365.25
    return float(ca._slice(to, start, end).iloc[1:].sum() / years)


def full_metrics(r: pd.Series, b: pd.Series) -> dict:
    '''Metrics of the whole window; r, b: daily log returns of run and benchmark.'''
    if len(r) < 2:
        return {}
    lvl = ca._level(r)
    cum, b_cum = cumulative_return(r), cumulative_return(b) if len(b) else np.nan
    return {
        'full_cum_return':       cum,
        'full_excess_return':    _excess(cum, b_cum),
        'full_max_dd':           maximum_drawdown(lvl),
        'full_ulcer':            ulcer_index(lvl),
        'full_ann_vol':          annualized_volatility(r),
        'full_underwater_share': float((lvl < lvl.cummax()).iloc[1:].mean()),
        'full_recovered':        bool(lvl.iloc[-1] >= lvl.iloc[0]),
    }


def overall_metrics(run: dict, r: pd.Series, start, end) -> dict:
    '''Extra full-period metrics of the overall sheet.'''
    if len(r) < 2:
        return {}
    return {
        'full_cagr':        annualized_return(r),
        'full_sharpe':      sharpe_ratio(r),
        'full_sortino':     sortino_ratio(r),
        'full_calmar':      calmar_ratio(r, ca._level(r)),
        'full_turnover_pa': _turnover_pa(run, r, start, end),
    }


def phase_metrics(r: pd.Series, b: pd.Series, phase: str) -> dict:
    '''PTT or TTP metrics; capture = run cum return / benchmark cum return.'''
    if len(r) == 0:
        return {}
    cum, b_cum = cumulative_return(r), cumulative_return(b) if len(b) else np.nan
    capture = 'down_capture' if phase == 'ptt' else 'recovery_capture'
    return {
        f'{phase}_cum_return':    cum,
        f'{phase}_excess_return': _excess(cum, b_cum),
        f'{phase}_{capture}':     _ratio(cum, b_cum),
        f'{phase}_max_dd':        maximum_drawdown(ca._level(r)),
    }


def window_row(run: dict, bench: pd.Series, w: dict, phases: bool, keys=KEY_COLS,
               info: dict | None = None) -> dict:
    '''One sheet row; info: per-group columns (role, base_run) taking precedence over run.'''
    info = info or {}
    row = {k: info[k] if k in info else run[k] for k in keys}
    s = lambda x, start, end: ca._slice(x, start, end)
    r = s(run['returns'], w['peak'], w['even'])
    row.update(full_metrics(r, s(bench, w['peak'], w['even'])))
    if not phases:                                    # overall sheet
        row.update(overall_metrics(run, r, w['peak'], w['even']))
    else:
        row.update(phase_metrics(s(run['returns'], w['peak'], w['trough']),
                                 s(bench, w['peak'], w['trough']), 'ptt'))
        row.update(phase_metrics(s(run['returns'], w['trough'], w['even']),
                                 s(bench, w['trough'], w['even']), 'ttp'))
    return row


# ----
# Regime runs: groups and detector diagnostics
# ----

def _regime_order(r: dict):
    return (ca.DETECTORS.index(r['detector']), r['series'], r['crisis_target'])


def _bench(runs: list[dict], freq: str) -> dict | None:
    return next((r for r in runs if r['model'] == BENCHMARK and r['frequency'] == freq), None)


def regime_groups(runs: list[dict]):
    '''
    (model, frequency, group, info) for every model of REGIME_MODELS and
    frequency: group = [benchmark, baseline, fallbacks, regime runs], the
    regime runs sorted by detector, series and crisis target.  Fallback = the
    portfolio a crisis target switches into, held all the time (FALLBACKS),
    for every target the model's regime runs use and that exists as a run.
    info: per run {'role', 'base_run'}; base_run = the model's baseline
    (empty for the benchmark).
    '''
    ca._baselines(runs)
    for freq in FREQ_ORDER:
        bench = _bench(runs, freq)
        if bench is None:
            continue
        for model in REGIME_MODELS:
            regimes = sorted((r for r in runs if r['model'] == model and r['frequency'] == freq
                              and r['detector'] != 'none' and r.get('base') is not None),
                             key=_regime_order)
            if not regimes:
                print(f'[crisis_analysis] no regime runs for {model} {freq}: skipped')
                continue
            base = regimes[0]['base']
            if any(r['base'] is not base for r in regimes):
                print(f'[crisis_analysis] {model} {freq}: several baselines, using {base["run"]}')
                regimes = [r for r in regimes if r['base'] is base]
            fallbacks = []
            for target in sorted({r['crisis_target'] for r in regimes}):
                if target not in FALLBACKS:
                    continue
                fb = _fallback(runs, target, freq)
                if fb is None:
                    print(f'[crisis_analysis] {model} {freq}: no fallback run for target {target}')
                elif fb is not base:
                    fallbacks.append(fb)
            group = [bench, base] + fallbacks + regimes
            roles = ['benchmark', 'baseline'] + ['fallback'] * len(fallbacks) + ['regime'] * len(regimes)
            info = [{'role': role, 'base_run': None if role == 'benchmark' else base['run'],
                     'fallback_run': _fallback_name(runs, r) if role == 'regime' else None}
                    for r, role in zip(group, roles)]
            yield model, freq, group, info


def _fallback(runs: list[dict], target: str, freq: str) -> dict | None:
    '''The run holding the fallback of crisis target `target` all the time, or None.'''
    test = FALLBACKS.get(target)
    return next((r for r in runs if r['frequency'] == freq and test(r)), None) if test else None


def _fallback_name(runs: list[dict], r: dict) -> str | None:
    fb = _fallback(runs, r['crisis_target'], r['frequency'])
    return fb['run'] if fb is not None else None


def all_groups(runs: list[dict]):
    '''
    (frequency, group, info) with every run of the frequency: benchmark,
    then per model (MODEL_ORDER) its baseline, controls and regime runs.
    Roles: benchmark, baseline, control (always min_var), regime;
    base_run = the run's own baseline (baselines: themselves).
    '''
    ca._baselines(runs)
    for freq in FREQ_ORDER:
        bench = _bench(runs, freq)
        if bench is None:
            continue
        others = [r for r in runs if r['frequency'] == freq and r is not bench]
        others.sort(key=lambda r: (_sort_key(r)[:2], r['detector'] != 'none',
                                   _regime_order(r) if r['detector'] != 'none' else (), r['run']))
        group = [bench] + others
        info = []
        for r in group:
            own_base = r['base']['run'] if r.get('base') else None
            if r is bench:
                info.append({'role': 'benchmark', 'base_run': None, 'fallback_run': None})
            elif r['detector'] != 'none':
                info.append({'role': 'regime', 'base_run': own_base,
                             'fallback_run': _fallback_name(runs, r)})
            elif r['crisis_target'] == 'always min_var':
                info.append({'role': 'control', 'base_run': own_base, 'fallback_run': None})
            else:
                info.append({'role': 'baseline', 'base_run': r['run'], 'fallback_run': None})
        yield freq, group, info


def defensive(run: dict) -> pd.Series | None:
    '''
    True on the trading days the run holds its crisis portfolio: the crisis
    flag of regime.csv, set on the rebalancing days, carried forward and
    applied from the next day on (the holding decided at a close earns the
    following returns).  None without regime.csv.
    '''
    if run['crisis'] is None:
        return None
    idx = run['returns'].index
    return run['crisis'].reindex(idx, method='ffill').shift(1).fillna(0.0) > 0.5


def detector_diagnostics(run: dict, c: dict) -> dict:
    '''
    Timing of the switch to the crisis portfolio in crisis c (defined dates),
    in trading days:
      exit_lag         defined peak -> first defensive day up to the trough;
                       negative = already defensive at the peak (the spell
                       began that many days earlier); empty = never defensive
                       peak -> trough
      reentry_lag      defined trough -> first day back in the model
                       portfolio; negative = back before the trough; empty =
                       not defensive peak -> trough, or never back
      days_defensive_before_peak
                       defensive at the peak: days since the switch (= -exit_lag);
                       otherwise 0
      late_defensive   True if not defensive peak -> trough but defensive
                       between trough and recovery: took the full fall, then
                       sat out (part of) the rebound
      defensive_after_trough_days
                       late_defensive only: trough -> first defensive day
      defensive_share  share of days defensive, peak -> recovery
      n_switches       changes between model and crisis portfolio, peak -> recovery
    _switch: date of the switch into the spell running at the peak (None if
    not defensive at the peak), for the cost of being defensive early.
    '''
    empty = {**dict.fromkeys(TIMING_COLS, np.nan), '_switch': None}
    d = defensive(run)
    if d is None:
        return empty
    on = d.to_numpy()
    pos = lambda date: int(d.index.searchsorted(pd.Timestamp(date), side='right')) - 1
    i_p, i_t, i_e = pos(c['peak']), pos(c['trough']), pos(c['even'])
    if i_p < 0:
        return empty

    switch = None
    ptt = on[i_p + 1:i_t + 1]
    if on[i_p]:
        j = i_p
        while j > 0 and on[j - 1]:
            j -= 1
        exit_lag, switch = j - i_p, d.index[j]
    elif ptt.any():
        exit_lag = int(np.argmax(ptt)) + 1
    else:
        exit_lag = np.nan

    if on[i_t]:
        rest = ~on[i_t + 1:]
        reentry_lag = int(np.argmax(rest)) + 1 if rest.any() else np.nan
    elif on[i_p:i_t + 1].any():
        last = i_p + int(np.flatnonzero(on[i_p:i_t + 1])[-1])
        reentry_lag = last + 1 - i_t
    else:
        reentry_lag = np.nan

    ttp = on[i_t + 1:i_e + 1]
    late = bool(not on[i_p:i_t + 1].any() and ttp.any())
    span = on[i_p:i_e + 1]
    return {
        'exit_lag':                    exit_lag,
        'reentry_lag':                 reentry_lag,
        'days_defensive_before_peak':  -exit_lag if switch is not None else 0,
        'late_defensive':              late,
        'defensive_after_trough_days': int(np.argmax(ttp)) + 1 if late else np.nan,
        'defensive_share':             float(span[1:].mean()) if len(span) > 1 else np.nan,
        'n_switches':                  int(np.abs(np.diff(span.astype(int))).sum()),
        '_switch':                     switch,
    }


def timing_table(group: list[dict], info: list[dict]) -> pd.DataFrame:
    '''
    One row per detector, series, frequency and crisis (TIMING_KEYS): the
    detector reads the market, not the portfolio, so the timing is the same
    for every model and crisis target.  n_runs = regime runs behind the row;
    identical = whether they all gave the same values (they should).
    '''
    rows = []
    for r, i in zip(group, info):
        if i['role'] != 'regime':
            continue
        for k, c in enumerate(main_crises, 1):
            dg = detector_diagnostics(r, c)
            rows.append({'frequency': r['frequency'], 'detector': r['detector'], 'series': r['series'],
                         'crisis': f'C{k}', **{col: dg[col] for col in TIMING_COLS}})
    if not rows:
        return pd.DataFrame(columns=TIMING_KEYS + TIMING_COLS)
    df = pd.DataFrame(rows)
    g = df.groupby(TIMING_KEYS, sort=False)
    out = g[TIMING_COLS].first()
    out['n_runs'] = g.size()
    out['identical'] = g[TIMING_COLS].nunique(dropna=False).max(axis=1).eq(1)
    return out.reset_index()


def add_switch_cost(sheets: dict, group: list[dict], info: list[dict]) -> None:
    '''
    Crisis sheets, regime runs already defensive at the peak: the cost of the
    early switch, from the switch date to the defined peak.
      missed_rally     the base model's return over that interval
      pre_peak_d_base  the run against its base model over that interval,
                       relative; with the cash target about -missed_rally
    Empty for every other row.
    '''
    by_name = {r['run']: r for r in group}
    for name, c, phases in WINDOWS:
        if not phases:
            continue
        missed, rel = [], []
        for r, i in zip(group, info):
            sw = detector_diagnostics(r, c)['_switch'] if i['role'] == 'regime' else None
            base = by_name.get(i.get('base_run'))
            if sw is None or base is None:
                missed.append(np.nan)
                rel.append(np.nan)
                continue
            b = cumulative_return(ca._slice(base['returns'], sw, c['peak']))
            o = cumulative_return(ca._slice(r['returns'], sw, c['peak']))
            missed.append(b)
            rel.append((1 + o) / (1 + b) - 1)
        sheets[name]['missed_rally'] = missed
        sheets[name]['pre_peak_d_base'] = rel


def _base_value(df: pd.DataFrame, col: str, ref: str = 'base_run') -> pd.Series:
    '''Value of col in the row named in each row's ref column (NaN without one).'''
    return df[ref].map(df.set_index('run')[col])


REFS = (('base_run', 'base'), ('fallback_run', 'fb'))      # (reference column, suffix)


def add_base_deltas(sheets: dict) -> None:
    '''
    <phase>_d_base / <phase>_d_fb: relative return against the row's base_run /
    fallback_run; empty without one.
    '''
    for df in sheets.values():
        for ref, suf in REFS:
            for phase in ('full', 'ptt', 'ttp'):
                col = f'{phase}_cum_return'
                if col in df and ref in df:
                    df[f'{phase}_d_{suf}'] = (1 + df[col]) / (1 + _base_value(df, col, ref)) - 1


def detector_averages(sheets_by_freq: dict) -> pd.DataFrame:
    '''
    Deltas against base and fallback averaged over the models, one row per
    frequency, crisis, detector, series and crisis target (n_models = runs
    averaged).  The models of one detector share its switching dates, so they
    are not independent; test these averages (or cluster by detector x
    frequency) instead of the model rows.
    '''
    parts = []
    for freq, sheets in sheets_by_freq.items():
        for name, df in sheets.items():
            reg = df[df['role'] == 'regime']
            cols = [c for c in reg.columns
                    if ('_d_base' in c or '_d_fb' in c) and reg[c].dtype.kind in 'fi']
            if reg.empty or not cols:
                continue
            g = reg.groupby(['detector', 'series', 'crisis_target'], sort=False)
            avg = g[cols].mean()
            avg.insert(0, 'n_models', g.size())
            avg = avg.reset_index()
            avg.insert(0, 'crisis', name)
            avg.insert(0, 'frequency', freq)
            parts.append(avg)
    return pd.concat(parts, ignore_index=True) if parts else pd.DataFrame()


# ----
# Main
# ----

def add_hit_rates(sheets: dict, tests: dict, ref_col: str | None = None, suffix: str = '') -> None:
    '''
    hit_rate<suffix>_<phase> on the overall sheet: share of the crisis sheets
    in which the run beats its reference row, test(row, ref_row) -> True /
    False, or NaN when not comparable (left out of the share);
    hit_count<suffix>_<phase> = the same as "hits/comparable crises", e.g. 2/5.
    Reference: the benchmark (ref_col None) or, per row, the run named in
    column ref_col.  Empty for the reference itself and rows without one.
    '''
    crises = [df for name, df in sheets.items() if name != 'overall']
    ov = sheets['overall']

    def refs(df):
        if ref_col is None:
            return pd.Series(df.loc[df['model'] == BENCHMARK, 'run'].iloc[0], index=df.index)
        return df[ref_col]

    ov_ref = refs(ov)
    empty = ov_ref.isna() | (ov_ref == ov['run'])
    for phase, test in tests.items():
        hits = []
        for df in crises:
            by_run = df.set_index('run', drop=False)
            h = [float(test(row, by_run.loc[ref])) if pd.notna(ref) and ref in by_run.index else np.nan
                 for (_, row), ref in zip(df.iterrows(), refs(df))]
            hits.append(pd.Series(h, index=df['run'].to_numpy()))
        h = pd.concat(hits, axis=1)
        count = h.sum(axis=1).astype(int).astype(str) + '/' + h.notna().sum(axis=1).astype(str)
        ov[f'hit_rate{suffix}_{phase}'] = ov['run'].map(h.mean(axis=1)).where(~empty)
        ov[f'hit_count{suffix}_{phase}'] = ov['run'].map(count).where(~empty)


# Beats the reference (benchmark or baseline): higher return over the phase
# (against the benchmark the same as excess return > 0)
HIT_TESTS = {
    'full': lambda r, b: r['full_cum_return'] > b['full_cum_return'],
    'ptt':  lambda r, b: r['ptt_cum_return'] > b['ptt_cum_return'],
    'ttp':  lambda r, b: r['ttp_cum_return'] > b['ttp_cum_return'],
}


def _sort_key(run: dict):
    m = run['model']
    return (MODEL_ORDER.index(m) if m in MODEL_ORDER else len(MODEL_ORDER), m, run['run'])


def _write(sheets: dict, path, n_keys: int = len(KEY_COLS)) -> None:
    from openpyxl.utils import get_column_letter
    with pd.ExcelWriter(path, engine='openpyxl') as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            ws.freeze_panes = get_column_letter(n_keys + 1) + '2'
            for j, col in enumerate(df.columns, 1):
                letter = get_column_letter(j)
                ws.column_dimensions[letter].width = min(max(10, len(str(col)) + 2), 32)
                fmt = '0.00%' if col in PCT_COLS else '0.000' if df[col].dtype.kind == 'f' else None
                if fmt:
                    for cell in ws[letter][1:]:
                        cell.number_format = fmt


def build_sheets(group: list[dict], bench: dict, keys=KEY_COLS, info=None) -> dict:
    info = info or [None] * len(group)
    sheets = {name: pd.DataFrame([window_row(r, bench['returns'], w, phases, keys, i)
                                  for r, i in zip(group, info)])
              for name, w, phases in WINDOWS}
    add_hit_rates(sheets, HIT_TESTS)
    return sheets


def run_all(tag: str, prefix: str, build, add_deltas, tests, write) -> None:
    '''
    The three outputs of both crisis analyses:
      1. baselines, one workbook per frequency       <prefix>_<frequency>.xlsx
      2. per regime model and frequency (REGIME_DIR) <prefix>_<model>_<frequency>.xlsx
      3. all runs, one workbook per frequency        <prefix>_all_<frequency>.xlsx
    build(group, bench, keys, info) -> sheets; add_deltas(sheets) adds the
    columns against base_run; tests = hit tests; write(sheets, path, n_keys).
    '''
    all_runs = [r for r in ca.find_runs() if r['model'] not in SKIP_MODELS]
    runs = [r for r in all_runs if r['detector'] == 'none']
    SAVE_DIR.mkdir(parents=True, exist_ok=True)
    for freq in FREQ_ORDER:
        group = sorted((r for r in runs if r['frequency'] == freq), key=_sort_key)
        bench = _bench(group, freq)
        if bench is None:
            print(f'[{tag}] no {BENCHMARK} run for {freq}: skipped')
            continue
        path = SAVE_DIR / f'{prefix}_{freq}.xlsx'
        write(build(group, bench), path, len(KEY_COLS))
        print(f'[{tag}] {freq}: {len(group)} portfolios, benchmark {bench["run"]}; saved to {path}')

    def regime_sheets(group, info):
        sheets = build(group, group[0], REGIME_KEYS, info)
        add_deltas(sheets)
        add_switch_cost(sheets, group, info)
        add_hit_rates(sheets, tests, ref_col='base_run', suffix='_base')
        add_hit_rates(sheets, tests, ref_col='fallback_run', suffix='_fb')
        return sheets

    REGIME_DIR.mkdir(parents=True, exist_ok=True)
    for model, freq, group, info in regime_groups(all_runs):
        path = REGIME_DIR / f'{prefix}_{model}_{freq}.xlsx'
        sheets = regime_sheets(group, info)
        sheets['timing'] = timing_table(group, info)
        write(sheets, path, len(REGIME_KEYS))
        n_fb = sum(i['role'] == 'fallback' for i in info)
        print(f'[{tag}] {model} {freq}: baseline {group[1]["run"]}, {n_fb} fallback, '
              f'{len(group) - 2 - n_fb} regime runs; saved to {path}')

    for freq, group, info in all_groups(all_runs):
        sheets = regime_sheets(group, info)
        avg = detector_averages({freq: sheets})
        sheets['timing'] = timing_table(group, info)
        sheets['detector_avg'] = avg
        if not sheets['timing']['identical'].all():
            print(f'[{tag}] {freq}: timing differs between models of the same detector: see column identical')
        path = SAVE_DIR / f'{prefix}_all_{freq}.xlsx'
        write(sheets, path, len(REGIME_KEYS))
        print(f'[{tag}] all runs {freq}: {len(group)} runs; saved to {path}')


def main() -> None:
    run_all('crisis_analysis', 'crisis_analysis', build_sheets, add_base_deltas, HIT_TESTS, _write)


if __name__ == '__main__':
    main()
