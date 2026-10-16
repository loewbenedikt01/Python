'''
Timing skill and detection lag of the regime detectors.

Sheets of timing.xlsx (in _output/_comparison/):

  lags        one row per detector, frequency and main crisis.  'daily' is the
              detector's own daily signal (p_crisis > THRESHOLD, regime file
              of _regimes/<detector>/regimes_final/); monthly / quarterly /
              yearly are the decisions a regime run actually took on its
              rebalance dates (raw/regime.csv of the first cash run found for
              that detector and frequency), so they include the waiting time
              until the next rebalance.
                entry_date      first day of the crisis spell that is on
                                between peak - LEAD_DAYS and the trough
                                (spells and gaps shorter than MIN_SPELL days
                                are ignored for entry and exit dates)
                already_on      the spell had started before peak - LEAD_DAYS
                entry_lag       trading days peak -> entry_date (negative:
                                before the peak)
                mkt_at_entry    market return peak -> entry_date (0 if before
                                the peak): the loss taken before protection
                ptt_share       share of peak -> trough days protected
                ptt_protected   share of the market's peak -> trough log loss
                                that fell on protected days
                on_at_trough    protected on the trough day
                exit_date       first day after the trough outside a spell
                                (only when on_at_trough; empty if still on at
                                the recovery date)
                exit_lag        trading days trough -> exit_date
                mkt_to_exit     market return trough -> exit_date: the rebound
                                missed before re-entering
                ttp_share       share of trough -> recovery days protected
                ttp_missed      share of the market's trough -> recovery log
                                gain that fell on protected days
              A day counts as protected when the decision in force at the
              previous close was crisis (the signal at d uses data through
              d - 1, the trade happens at the close of d, so the first
              protected return is that of d + 1).
  lag_summary per detector and frequency: median entry and exit lag and
              the other columns averaged over the crises
  classify    the daily protection flag against three definitions of a stress
              day over the full period (crises.overall): inside a main crisis
              peak -> trough window; market drawdown from its running peak of
              at least STRESS_DD; realised market volatility over the next
              VOL_WINDOW days of at least STRESS_VOL.  share_on, hit rate
              (recall), precision, false alarm rate (flagged share of
              non-stress days), balanced accuracy and Cohen's kappa.
  agreement   pairwise agreement of the daily detector flags: share of days
              with the same label, Cohen's kappa, correlation of p_crisis
  timing_reg  market-timing regressions of every run on the market, monthly
              excess returns (simple, over the T-bill):
                Henriksson-Merton  r_p = a + b r_m + c max(0, -r_m) + e
                Treynor-Mazuy      r_p = a + b r_m + g r_m^2 + e
              c > 0 (g > 0): lower beta in down months than in up months,
              i.e. timing skill.  beta_down = b - c.  Newey-West standard
              errors (HAC_LAGS).
  timing_summary  mean c and g and the number of runs with c (g) significantly
              above 0 at 5 %, per detector, crisis target and frequency

The market is the model folder MARKET_MODEL (the S&P 500 run) or, if it does
not exist, compare_all.BENCH_MODEL.

Usage: set the variables below and run the file.
'''

from pathlib import Path

import numpy as np
import pandas as pd
from scipy import stats

import compare_all as ca
from compare_runs import CRISIS_NAMES
from crises import overall, main_crises
from risk_free_rate import rf_daily

# ----
# Variables
# ----

DETECTORS    = {'changepoint': 'vix', 'hmm': 'gspc', 'wasserstein': 'gspc_vix'}   # detector -> series
THRESHOLD    = 0.5          # p_crisis above this = crisis (regime_def.CRISIS_THRESHOLD)
MARKET_MODEL = 'sp500'      # model folder of the market series
LEAD_DAYS    = 63           # a spell starting up to this many trading days before the peak counts
MIN_SPELL    = 5            # entry / exit dates: spells and gaps shorter than this many days are ignored
STRESS_DD    = 0.10         # classify: market drawdown that marks a stress day
STRESS_VOL   = 0.20         # classify: forward realised volatility that marks a stress day
VOL_WINDOW   = 21           # classify: trading days of the forward volatility
HAC_LAGS     = None         # timing_reg: Newey-West lags; None = floor(4 (T / 100)^(2 / 9))
ALPHA        = 0.05

ROOT     = Path(__file__).resolve().parents[1]
SAVE_DIR = ca.SAVE_DIR
FREQS    = ca.FREQ_ORDER


# ----
# Data
# ----

def market_returns(runs: list[dict]) -> pd.Series:
    '''Daily log returns of the market (MARKET_MODEL, else BENCH_MODEL).'''
    for name in (MARKET_MODEL, ca.BENCH_MODEL):
        cand = [r for r in runs if r['model'] == name]
        if cand:
            if name != MARKET_MODEL:
                label = name + '/' + cand[0]['run']
                print(f'[timing] no {MARKET_MODEL} folder: market = {label}')
            return cand[0]['returns']
    raise FileNotFoundError(f'[timing] neither {MARKET_MODEL} nor {ca.BENCH_MODEL} in _output')


def daily_signal(detector: str, series: str) -> pd.Series | None:
    '''p_crisis of the detector file (value at d uses data through d - 1).'''
    path = ROOT / '_regimes' / detector / 'regimes_final' / f'{detector}_{series}.csv'
    if not path.exists():
        print(f'[timing] {path} not found: {detector} skipped')
        return None
    df = pd.read_csv(path, index_col='date', parse_dates=['date']).sort_index()
    return df['p_crisis'].astype(float)


def decisions(runs: list[dict], detector: str, series: str, freq: str) -> pd.Series | None:
    '''crisis flag (True/False) on the rebalance dates of one regime run.'''
    cand = [r for r in runs if r['detector'] == detector and r['series'] == series
            and r['frequency'] == freq and r['crisis'] is not None]
    cand.sort(key=lambda r: r['crisis_target'] != 'cash')          # cash runs first
    return cand[0]['crisis'].astype(bool) if cand else None


def protected(flag: pd.Series, idx: pd.DatetimeIndex) -> pd.Series:
    '''
    True on the trading days of idx whose return is earned in the crisis
    portfolio: the decision in force at the previous trading day's close.
    flag: True/False on decision (or signal) dates.
    '''
    state = flag.astype(float).reindex(idx.union(flag.index)).ffill().reindex(idx)
    return state.shift(1).fillna(0.0).astype(bool)


# ----
# Lags around the crises
# ----

def _pos(idx: pd.DatetimeIndex, d) -> int:
    '''Position of the first trading day on or after d.'''
    return int(idx.searchsorted(pd.Timestamp(d)))


def _spells(on: np.ndarray, value: bool) -> list[tuple[int, int]]:
    '''(start, end) positions, end inclusive, of the runs of on == value.'''
    x = np.concatenate([[False], on == value, [False]]).astype(int)
    d = np.diff(x)
    return list(zip(np.flatnonzero(d == 1), np.flatnonzero(d == -1) - 1))


def steady(on: np.ndarray, min_len: int = MIN_SPELL) -> np.ndarray:
    '''
    The protection flag without flicker, for the entry and exit dates only:
    off gaps shorter than min_len inside protection are filled, then
    protected spells shorter than min_len are dropped.
    '''
    out = on.copy()
    for a, b in _spells(out, False):
        if b - a + 1 < min_len and a > 0 and b < len(out) - 1:
            out[a:b + 1] = True
    for a, b in _spells(out, True):
        if b - a + 1 < min_len:
            out[a:b + 1] = False
    return out


def crisis_lags(prot: pd.Series, mkt: pd.Series, c: dict) -> dict:
    idx = mkt.index
    p0, p1, p2 = _pos(idx, c['peak']), _pos(idx, c['trough']), _pos(idx, c['even'])
    p2 = min(p2, len(idx) - 1)
    lvl = np.exp(mkt.cumsum())
    out = {'peak': idx[p0].date(), 'trough': idx[p1].date(), 'recovery': idx[p2].date()}
    on = prot.to_numpy()
    st = steady(on)

    # entry: start of the first steady spell that is on in [peak - LEAD_DAYS, trough]
    lo = max(p0 - LEAD_DAYS, 0)
    window = np.flatnonzero(st[lo:p1 + 1])
    if len(window):
        e = lo + int(window[0])
        out['already_on'] = bool(e == lo and lo > 0 and st[lo - 1])
        while e > 0 and st[e - 1]:
            e -= 1
        out['entry_date'] = idx[e].date()
        out['entry_lag'] = e - p0
        out['mkt_at_entry'] = float(lvl.iloc[e - 1] / lvl.iloc[p0] - 1.0) if e > p0 + 1 else 0.0
    else:
        out.update({'already_on': False, 'entry_date': None, 'entry_lag': np.nan, 'mkt_at_entry': np.nan})

    ptt = slice(p0 + 1, p1 + 1)                     # returns in (peak, trough]
    r, f = mkt.iloc[ptt].to_numpy(), on[ptt]
    out['ptt_share'] = float(f.mean()) if len(f) else np.nan
    out['ptt_protected'] = float(r[f].sum() / r.sum()) if r.sum() < 0 else np.nan

    # exit: first day after the trough outside a steady spell, if protected at the trough
    after = np.flatnonzero(~st[p1 + 1:p2 + 1])
    out['on_at_trough'] = bool(st[p1])
    if st[p1] and len(after):
        x = p1 + 1 + int(after[0])
        out['exit_date'] = idx[x].date()
        out['exit_lag'] = x - p1
        out['mkt_to_exit'] = float(lvl.iloc[x - 1] / lvl.iloc[p1] - 1.0)
    else:
        out.update({'exit_date': None, 'exit_lag': np.nan, 'mkt_to_exit': np.nan})

    ttp = slice(p1 + 1, p2 + 1)
    r, f = mkt.iloc[ttp].to_numpy(), on[ttp]
    out['ttp_share'] = float(f.mean()) if len(f) else np.nan
    out['ttp_missed'] = float(r[f].sum() / r.sum()) if r.sum() > 0 else np.nan
    return out


def lag_table(flags: dict, mkt: pd.Series) -> pd.DataFrame:
    rows = []
    for (det, freq), prot in flags.items():
        for i, c in enumerate(main_crises, 1):
            name = f'C{i} ' + CRISIS_NAMES.get(c['label'], c['label'])
            rows.append({'detector': det, 'frequency': freq, 'crisis': name, **crisis_lags(prot, mkt, c)})
    return pd.DataFrame(rows)


def lag_summary(lags: pd.DataFrame) -> pd.DataFrame:
    '''Mean over the crises (median for the two lags, which a single crisis can dominate).'''
    cols = ['mkt_at_entry', 'ptt_share', 'ptt_protected', 'mkt_to_exit', 'ttp_share', 'ttp_missed']
    g = lags.groupby(['detector', 'frequency'], sort=False)
    out = pd.concat([g['entry_lag'].median().rename('entry_lag_median'),
                     g['exit_lag'].median().rename('exit_lag_median'), g[cols].mean()], axis=1)
    out['crises_entered'] = g['entry_date'].apply(lambda s: int(s.notna().sum()))
    out['already_on'] = g['already_on'].sum().astype(int)
    out['on_at_trough'] = g['on_at_trough'].sum().astype(int)
    return out.reset_index()


# ----
# Classification against stress days
# ----

def stress_days(mkt: pd.Series) -> dict[str, pd.Series]:
    lvl = np.exp(mkt.cumsum())
    dd = 1.0 - lvl / lvl.cummax()
    in_ptt = pd.Series(False, index=mkt.index)
    for c in main_crises:
        in_ptt |= (mkt.index > pd.Timestamp(c['peak'])) & (mkt.index <= pd.Timestamp(c['trough']))
    fwd_vol = (mkt.rolling(VOL_WINDOW).std() * np.sqrt(ca.TRADING_DAYS_PER_YEAR)).shift(-(VOL_WINDOW - 1))
    return {'crisis peak -> trough': in_ptt,
            f'drawdown >= {STRESS_DD:.0%}': dd >= STRESS_DD,
            f'forward {VOL_WINDOW}d vol >= {STRESS_VOL:.0%}': (fwd_vol >= STRESS_VOL).where(fwd_vol.notna())}


def _kappa(a: np.ndarray, b: np.ndarray) -> float:
    po = (a == b).mean()
    pe = a.mean() * b.mean() + (1 - a.mean()) * (1 - b.mean())
    return float((po - pe) / (1 - pe)) if pe < 1 else np.nan


def classify(flags: dict, mkt: pd.Series) -> pd.DataFrame:
    truth = stress_days(mkt)
    w = overall[0]
    rows = []
    for (det, freq), prot in flags.items():
        for name, t in truth.items():
            df = pd.concat([prot, t], axis=1, keys=['f', 't']).dropna()
            df = df.loc[(df.index > pd.Timestamp(w['peak'])) & (df.index <= pd.Timestamp(w['even']))]
            f, t = df['f'].astype(bool).to_numpy(), df['t'].astype(bool).to_numpy()
            tp, fp = (f & t).sum(), (f & ~t).sum()
            fn, tn = (~f & t).sum(), (~f & ~t).sum()
            hit = tp / (tp + fn) if tp + fn else np.nan
            fa = fp / (fp + tn) if fp + tn else np.nan
            rows.append({'detector': det, 'frequency': freq, 'stress definition': name,
                         'days': len(f), 'stress_share': t.mean(), 'share_on': f.mean(),
                         'hit_rate': hit, 'precision': tp / (tp + fp) if tp + fp else np.nan,
                         'false_alarm_rate': fa, 'balanced_accuracy': (hit + 1 - fa) / 2,
                         'kappa': _kappa(f, t)})
    return pd.DataFrame(rows)


def agreement(signals: dict[str, pd.Series]) -> pd.DataFrame:
    w = overall[0]
    rows = []
    names = list(signals)
    for i, a in enumerate(names):
        for b in names[i + 1:]:
            df = pd.concat([signals[a], signals[b]], axis=1, keys=['a', 'b'], sort=True).dropna()
            df = df.loc[(df.index > pd.Timestamp(w['peak'])) & (df.index <= pd.Timestamp(w['even']))]
            fa, fb = (df['a'] > THRESHOLD).to_numpy(), (df['b'] > THRESHOLD).to_numpy()
            rows.append({'detector a': a, 'detector b': b, 'days': len(df),
                         'same_label': float((fa == fb).mean()), 'kappa': _kappa(fa, fb),
                         'corr_p_crisis': float(df['a'].corr(df['b']))})
    return pd.DataFrame(rows)


# ----
# Timing regressions
# ----

def _monthly_excess(r: pd.Series) -> pd.Series:
    '''Monthly simple excess return over the T-bill from daily log returns.'''
    df = pd.DataFrame({'r': r, 'rf': rf_daily(r.index).to_numpy()}, index=r.index)
    m = df.groupby(df.index.to_period('M')).sum()
    return np.expm1(m['r']) - np.expm1(m['rf'])


def ols_hac(y: np.ndarray, X: np.ndarray, lags: int | None = None) -> tuple[np.ndarray, np.ndarray]:
    '''OLS coefficients and Newey-West standard errors (Bartlett kernel).'''
    T = len(y)
    b, *_ = np.linalg.lstsq(X, y, rcond=None)
    u = y - X @ b
    L = int(np.floor(4 * (T / 100) ** (2 / 9))) if lags is None else lags
    g = X * u[:, None]
    S = g.T @ g
    for l in range(1, L + 1):
        G = g[l:].T @ g[:-l]
        S += (1 - l / (L + 1)) * (G + G.T)
    XtX_inv = np.linalg.inv(X.T @ X)
    cov = XtX_inv @ S @ XtX_inv
    return b, np.sqrt(np.diag(cov))


def timing_regressions(runs: list[dict], mkt: pd.Series) -> pd.DataFrame:
    keys = ['model', 'frequency', 'detector', 'series', 'crisis_target', 'run']
    m = _monthly_excess(mkt)
    rows, seen = [], set()
    for run in sorted(runs, key=ca._sort_key):
        if run['model'] == MARKET_MODEL or id(run['returns']) in seen:
            continue
        seen.add(id(run['returns']))
        p = _monthly_excess(run['returns'])
        df = pd.concat([p, m], axis=1, keys=['p', 'm'], join='inner').dropna()
        if len(df) < 24:
            continue
        y, x = df['p'].to_numpy(), df['m'].to_numpy()
        T = len(y)
        row = {k: run[k] for k in keys}
        row['months'] = T
        X = np.column_stack([np.ones(T), x, np.maximum(0.0, -x)])
        b, se = ols_hac(y, X, HAC_LAGS)
        t = b / se
        row.update({'hm_alpha_ann': b[0] * 12, 'hm_p_alpha': 2 * stats.t.sf(abs(t[0]), T - 3),
                    'beta_up': b[1], 'beta_down': b[1] - b[2],
                    'hm_c': b[2], 'hm_t_c': t[2], 'hm_p_c': 2 * stats.t.sf(abs(t[2]), T - 3)})
        X = np.column_stack([np.ones(T), x, x ** 2])
        b, se = ols_hac(y, X, HAC_LAGS)
        t = b / se
        row.update({'tm_alpha_ann': b[0] * 12, 'tm_g': b[2], 'tm_t_g': t[2],
                    'tm_p_g': 2 * stats.t.sf(abs(t[2]), T - 3)})
        rows.append(row)
    return pd.DataFrame(rows)


def timing_summary(reg: pd.DataFrame) -> pd.DataFrame:
    df = reg.copy()
    df['hm_sig_pos'] = (df['hm_c'] > 0) & (df['hm_p_c'] < ALPHA)
    df['hm_sig_neg'] = (df['hm_c'] < 0) & (df['hm_p_c'] < ALPHA)
    df['tm_sig_pos'] = (df['tm_g'] > 0) & (df['tm_p_g'] < ALPHA)
    df['crisis_target'] = df['crisis_target'].replace('', 'baseline')
    g = df.groupby(['detector', 'crisis_target', 'frequency'], sort=False)
    out = g.agg(runs=('run', 'size'), hm_c_mean=('hm_c', 'mean'), beta_up_mean=('beta_up', 'mean'),
                beta_down_mean=('beta_down', 'mean'), hm_sig_pos=('hm_sig_pos', 'sum'),
                hm_sig_neg=('hm_sig_neg', 'sum'), tm_g_mean=('tm_g', 'mean'),
                tm_sig_pos=('tm_sig_pos', 'sum'), hm_alpha_ann_mean=('hm_alpha_ann', 'mean'))
    return out.reset_index()


# ----
# Main
# ----

def _write(sheets: dict, path: Path) -> None:
    from openpyxl.utils import get_column_letter
    pct = {'mkt_at_entry', 'ptt_share', 'ptt_protected', 'mkt_to_exit', 'ttp_share', 'ttp_missed',
           'stress_share', 'share_on', 'hit_rate', 'precision', 'false_alarm_rate',
           'balanced_accuracy', 'same_label', 'hm_alpha_ann', 'tm_alpha_ann', 'hm_alpha_ann_mean'}
    with pd.ExcelWriter(path, engine='openpyxl') as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            ws.freeze_panes = 'B2'
            for j, col in enumerate(df.columns, 1):
                letter = get_column_letter(j)
                ws.column_dimensions[letter].width = min(max(10, len(str(col)) + 2), 40)
                fmt = '0.0%' if col in pct else '0.000' if df[col].dtype.kind == 'f' else None
                if fmt:
                    for cell in ws[letter][1:]:
                        cell.number_format = fmt


def main() -> None:
    runs = ca.find_runs()
    if not runs:
        print(f'[timing] no runs found in {ca.OUTPUT_ROOT}')
        return
    mkt = market_returns(runs)
    signals, flags = {}, {}
    for det, series in DETECTORS.items():
        s = daily_signal(det, series)
        if s is None:
            continue
        signals[det] = s
        flags[(det, 'daily')] = protected(s > THRESHOLD, mkt.index)
        for freq in FREQS:
            d = decisions(runs, det, series, freq)
            if d is not None:
                flags[(det, freq)] = protected(d, mkt.index)
    lags = lag_table(flags, mkt)
    reg = timing_regressions(runs, mkt)
    sheets = {'lags': lags, 'lag_summary': lag_summary(lags), 'classify': classify(flags, mkt),
              'agreement': agreement(signals), 'timing_reg': reg, 'timing_summary': timing_summary(reg)}
    SAVE_DIR.mkdir(parents=True, exist_ok=True)
    path = SAVE_DIR / 'timing.xlsx'
    _write(sheets, path)
    print(f'[timing] {len(flags)} signals, {len(reg)} runs regressed; saved to {path}')


if __name__ == '__main__':
    main()
